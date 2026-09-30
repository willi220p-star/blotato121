"""CSV, JSON, SQLite and Excel exports."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database.database import write_sqlite

COMPANY_FIELDS = [
    "company_id",
    "company_name",
    "trading_name",
    "legal_name",
    "abn",
    "acn",
    "description",
    "website",
    "industry",
    "services",
    "address",
    "suburb",
    "state",
    "postcode",
    "country",
    "phone",
    "phone_original",
    "phone_normalized",
    "email",
    "general_email",
    "accounts_email",
    "info_email",
    "contact_email",
    "website_contact_page",
    "linkedin_url",
    "facebook_url",
    "instagram_url",
    "source_url",
    "source_type",
    "date_found",
    "confidence_score",
    "verification_status",
    "scrape_status",
]
EMPLOYEE_FIELDS = [
    "company_id",
    "company_name",
    "first_name",
    "last_name",
    "full_name",
    "job_title",
    "seniority",
    "qualification",
    "location",
    "public_work_email",
    "email_type",
    "email_source_url",
    "business_phone",
    "linkedin_url",
    "social_url",
    "source_url",
    "source_type",
    "date_found",
]
SOURCE_FIELDS = ["company_id", "employee_id", "source_url", "source_type", "date_accessed", "information_found"]


def _write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field) if row.get(field) not in {None} else "" for field in fields})


def _xlsx(path: Path, companies: list[dict], employees: list[dict], sources: list[dict]) -> None:
    workbook = Workbook()
    how = workbook.active
    how.title = "How to use"
    how.append(["Darwin accounting firms — public research"])
    how.append(["Official websites + public directories only. No LinkedIn crawl. No guessed emails."])
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for title, rows, fields in (
        ("Companies", companies, COMPANY_FIELDS),
        ("Employees", employees, EMPLOYEE_FIELDS),
        ("Sources", sources, SOURCE_FIELDS),
    ):
        sheet = workbook.create_sheet(title)
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(field) if row.get(field) not in {None} else "" for field in fields])
        for index, header in enumerate(fields, start=1):
            cell = sheet.cell(1, index)
            cell.fill = fill
            cell.font = font
            sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 14), 40)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def completeness(companies: list[dict], employees: list[dict], sources: list[dict], extras: dict | None = None) -> dict:
    extras = extras or {}
    return {
        "companies_found": len(companies),
        "companies_verified": sum(1 for row in companies if (row.get("confidence_score") or 0) >= 75),
        "employees_found": len(employees),
        "public_business_emails": sum(1 for row in companies if row.get("email"))
        + sum(1 for row in employees if row.get("public_work_email")),
        "public_business_phones": sum(1 for row in companies if row.get("phone")),
        "sources_collected": len(sources),
        "duplicates_removed": extras.get("duplicates_removed", 0),
        "errors": extras.get("errors", 0),
        "companies_with_websites": sum(1 for row in companies if row.get("website")),
        "companies_without_public_staff": sum(
            1 for row in companies if not any(person.get("company_id") == row.get("company_id") for person in employees)
        ),
    }


def write_outputs(
    companies: list[dict],
    employees: list[dict],
    sources: list[dict],
    output_dir: Path,
    extra_dirs: list[Path] | None = None,
    extras: dict | None = None,
) -> dict:
    stats = completeness(companies, employees, sources, extras)
    output_dir.mkdir(parents=True, exist_ok=True)
    targets = [output_dir, *(extra_dirs or [])]
    payload = {"stats": stats, "companies": companies, "employees": employees, "sources": sources}
    for folder in targets:
        folder.mkdir(parents=True, exist_ok=True)
        _write_csv(folder / "darwin_accounting_companies.csv", companies, COMPANY_FIELDS)
        _write_csv(folder / "darwin_accounting_employees.csv", employees, EMPLOYEE_FIELDS)
        _write_csv(folder / "darwin_accounting_sources.csv", sources, SOURCE_FIELDS)
        (folder / "darwin_accounting.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        write_sqlite(folder / "darwin_accounting.db", companies, employees, sources)
        _xlsx(folder / "Darwin_accounting_firms.xlsx", companies, employees, sources)
    return stats
