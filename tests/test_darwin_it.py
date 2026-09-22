import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from lib.darwin_it import (
    classify_it_role,
    extract_job_openings,
    is_darwin_location,
    merge_companies,
    parse_ictnt_list,
    parse_ictnt_profile,
    plausible_it_job_title,
    rank_resume_email,
    write_workbook,
)

ICTNT_LIST = """
<html><body>
<div class="list-item">
  <a href="/ict-directory/219">view profile</a>
  <div class="left-column"><strong>ACTEK</strong><br>Darwin</div>
</div>
<div class="list-item">
  <a href="/ict-directory/223">view profile</a>
  <div class="left-column"><strong>Red Centre Technology Partners</strong><br>Alice-springs</div>
</div>
<div class="list-item">
  <a href="/ict-directory/150">view profile</a>
  <div class="left-column"><strong>Radical Systems</strong><br>Darwin</div>
</div>
</body></html>
"""

ICTNT_PROFILE = """
<html><body>
<h2>You are here</h2>
<h2>Radical Systems</h2>
<a href="https://www.radicalsystems.com.au/">visit company website</a>
<a href="mailto:directors@radicalsystems.com.au">directors@radicalsystems.com.au</a>
<a href="mailto:info@ictnt.asn.au">info@ictnt.asn.au</a>
<h3>Contact Us</h3>
<p>1300 757 155 Level 3/9-11 Cavenagh Street Darwin NT 0800</p>
</body></html>
"""

JOB_PAGE = """
<html><body>
<script type="application/ld+json">
{
  "@type":"JobPosting",
  "title":"Senior Systems Engineer",
  "url":"/jobs/senior-systems-engineer",
  "datePosted":"2026-09-01",
  "validThrough":"2026-12-31",
  "employmentType":"FULL_TIME",
  "jobLocation":{"address":{"addressLocality":"Darwin","addressRegion":"NT"}}
}
</script>
<script type="application/ld+json">
{"@type":"JobPosting","title":"Accountant","url":"/jobs/accountant"}
</script>
<script type="application/ld+json">
{
  "@type":"JobPosting",
  "title":"IT Support Officer",
  "url":"/jobs/old",
  "validThrough":"2020-01-01"
}
</script>
<h2>Careers</h2>
<a href="/careers/software-developer">Software Developer</a>
<p>Managed IT Services</p>
</body></html>
"""


class DarwinITTest(unittest.TestCase):
    def test_classifies_it_roles(self):
        cases = {
            "Senior Systems Engineer": "Systems / Infrastructure Engineer",
            "IT Support Officer": "IT Support / Helpdesk",
            "Full Stack Developer": "Software Developer / Engineer",
            "Cyber Security Analyst": "Cybersecurity",
            "ICT Manager": "ICT Manager / Team Lead",
            "Azure Cloud Engineer": "Cloud / Microsoft 365",
        }
        for title, expected in cases.items():
            with self.subTest(title=title):
                self.assertEqual(classify_it_role(title), expected)

    def test_rejects_service_labels_as_jobs(self):
        self.assertFalse(plausible_it_job_title("IT Support"))
        self.assertFalse(plausible_it_job_title("Managed IT Services"))
        self.assertFalse(plausible_it_job_title("What is a software engineer?"))
        self.assertTrue(plausible_it_job_title("IT Support Officer"))
        self.assertTrue(plausible_it_job_title("Software Developer"))

    def test_prefers_careers_email_over_gmail(self):
        ranked = rank_resume_email(
            ["hello@trueblueit.com.au", "jobs@trueblueit.com.au", "owner@gmail.com"],
            "https://trueblueit.com.au/",
        )
        self.assertEqual(ranked["best_resume_email"], "jobs@trueblueit.com.au")
        self.assertEqual(ranked["email_type"], "careers / HR inbox")
        self.assertEqual(ranked["gmail_email"], "owner@gmail.com")

    def test_uses_published_gmail_when_that_is_the_inbox(self):
        ranked = rank_resume_email(["fixitmichael@gmail.com"])
        self.assertEqual(ranked["best_resume_email"], "fixitmichael@gmail.com")
        self.assertEqual(ranked["email_type"], "published Gmail")

    def test_parses_darwin_ictnt_list_and_skips_alice_springs(self):
        companies = parse_ictnt_list(ICTNT_LIST)
        names = {item["name"]: item["location"] for item in companies}
        self.assertEqual(names["ACTEK"], "Darwin")
        self.assertEqual(names["Radical Systems"], "Darwin")
        self.assertEqual(names["Red Centre Technology Partners"], "Alice Springs")
        self.assertFalse(is_darwin_location("Alice Springs"))
        self.assertTrue(is_darwin_location("Winnellie, Darwin NT"))

    def test_parses_ictnt_profile_website_and_email(self):
        parsed = parse_ictnt_profile(ICTNT_PROFILE, "https://www.ictnt.asn.au/ict-directory/150")
        self.assertEqual(parsed["name"], "Radical Systems")
        self.assertEqual(parsed["website"], "https://www.radicalsystems.com.au/")
        self.assertEqual(parsed["directory_email"], "directors@radicalsystems.com.au")

    def test_extracts_current_it_job_postings(self):
        openings = extract_job_openings(JOB_PAGE, "https://example.com/careers")
        titles = {item["position_title"] for item in openings}
        self.assertIn("Senior Systems Engineer", titles)
        self.assertIn("Software Developer", titles)
        self.assertNotIn("Accountant", titles)
        self.assertNotIn("IT Support Officer", titles)

    def test_merges_directory_and_extra_companies(self):
        merged = merge_companies(
            [
                {
                    "name": "Radical Systems",
                    "location": "Darwin",
                    "website": "https://www.radicalsystems.com.au/",
                    "source": "ICTNT directory",
                },
                {
                    "name": "Red Centre Technology Partners",
                    "location": "Alice Springs",
                    "website": "https://example.net/",
                    "source": "ICTNT directory",
                },
            ],
            [
                {
                    "name": "Get A Geek NT",
                    "website": "https://getageeknt.com.au/",
                    "source": "Official website",
                    "suburb": "Darwin",
                    "employer_type": "Private IT company",
                }
            ],
        )
        names = [item["name"] for item in merged]
        self.assertIn("Radical Systems", names)
        self.assertIn("Get A Geek NT", names)
        self.assertNotIn("Red Centre Technology Partners", names)

    def test_writes_company_job_and_email_sheets(self):
        results = [
            {
                "name": "BlueReef Technology",
                "jobs_available": 1,
                "best_resume_email": "help@bluereef.tech",
                "email_type": "general company inbox",
                "gmail_email": "",
                "other_emails": "",
                "website": "https://bluereef.tech/",
                "suburb": "Winnellie",
                "employer_type": "Private IT company",
                "focus": "Managed IT",
                "phone": "08 8922 0000",
                "source": "Official website",
                "research_status": "jobs published",
                "profile_url": "",
                "pages_checked": ["https://bluereef.tech/careers"],
                "openings": [
                    {
                        "position_title": "IT Support Officer",
                        "role_category": "IT Support / Helpdesk",
                        "job_url": "https://bluereef.tech/careers",
                        "location": "Darwin",
                        "employment_type": "FULL_TIME",
                        "date_posted": "",
                        "valid_through": "",
                        "status": "Published careers-page listing; closing date not supplied",
                        "source_page": "https://bluereef.tech/careers",
                    }
                ],
            },
            {
                "name": "Quiet MSP",
                "jobs_available": 0,
                "best_resume_email": "info@quiet.example",
                "email_type": "general company inbox",
                "gmail_email": "",
                "other_emails": "",
                "website": "https://quiet.example/",
                "suburb": "Darwin",
                "employer_type": "Private IT company",
                "focus": "",
                "phone": "",
                "source": "Official website",
                "research_status": "official site checked; no published IT jobs extracted",
                "profile_url": "",
                "pages_checked": [],
                "openings": [],
            },
        ]
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "darwin.xlsx"
            summary = write_workbook(path, results, "2026-09-22T00:00:00+00:00")
            self.assertEqual(summary["companies"], 2)
            self.assertEqual(summary["hiring"], 1)
            self.assertEqual(summary["jobs"], 1)
            workbook = load_workbook(path)
            self.assertEqual(
                workbook.sheetnames,
                ["How to use", "Companies", "Hiring now", "Job listings", "Resume emails"],
            )
            self.assertEqual(workbook["Hiring now"].max_row, 2)
            self.assertEqual(workbook["Job listings"]["B2"].value, "IT Support Officer")
            self.assertEqual(workbook["Resume emails"]["B2"].value, "help@bluereef.tech")


if __name__ == "__main__":
    unittest.main()
