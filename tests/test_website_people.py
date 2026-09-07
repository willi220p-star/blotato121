import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from lib.website_people import (
    extract_people_from_html,
    merge_public_search_people,
    normalise_person_linkedin_url,
    organisations_from_workbook,
    write_enriched_workbook,
)


class WebsitePeopleTest(unittest.TestCase):
    def test_normalise_person_linkedin_url(self):
        self.assertEqual(
            normalise_person_linkedin_url(
                "https://au.linkedin.com/in/jane-smith-123/?trk=profile"
            ),
            "https://www.linkedin.com/in/jane-smith-123",
        )
        self.assertEqual(
            normalise_person_linkedin_url("https://linkedin.com/company/example"),
            "",
        )

    def test_extracts_official_link_and_structured_person(self):
        page = """
        <html><body>
          <section class="team-card">
            <h3>Jane Smith</h3>
            <p>Chief Executive Officer</p>
            <a href="https://www.linkedin.com/in/jane-smith-123">LinkedIn</a>
          </section>
          <script type="application/ld+json">
            {"@type":"Person","name":"David Jones","jobTitle":"Director"}
          </script>
          <p>Heidi Coleman – Head Of Finance</p>
          <h3>Tre Manning-Watson</h3>
          <p>Regional Manager – Northern Territory &amp; South Australia</p>
          <h3>Jack Hatcher</h3>
          <p>Regional Manager – Cairns</p>
        </body></html>
        """
        people = extract_people_from_html(page, "https://example.org/our-team")
        by_name = {person["person_name"]: person for person in people}
        self.assertEqual(
            by_name["Jane Smith"]["linkedin_profile_url"],
            "https://www.linkedin.com/in/jane-smith-123",
        )
        self.assertIn("Chief Executive Officer", by_name["Jane Smith"]["role"])
        self.assertEqual(by_name["David Jones"]["role"], "Director")
        self.assertEqual(by_name["Heidi Coleman"]["role"], "Head Of Finance")
        self.assertEqual(
            by_name["Tre Manning-Watson"]["role"],
            "Regional Manager – Northern Territory & South Australia",
        )
        self.assertEqual(by_name["Jack Hatcher"]["role"], "Regional Manager – Cairns")

    def test_excludes_article_headings_that_are_not_people(self):
        page = """
        <html><body>
          <article>
            <h2>How you pay for providers</h2>
            <p>Before your Plan Manager starts helping you manage your plan.</p>
            <h3>Meet Our Team</h3>
            <p>Our team includes experienced managers.</p>
            <h3>Support Coordination</h3>
            <p>Our Manager can explain the service.</p>
            <h3>Richard Drevid</h3>
            <h3>Alextina Javi Manager</h3>
            <h3>Board Members</h3>
            <p>Board Director</p>
            <h3>Voyage Support</h3>
            <p>Services Manager</p>
            <section>
              <h3>Mahesh Perera</h3>
              <p>Director</p>
              <a href="https://www.linkedin.com/in/freehold-lisasmith">LinkedIn</a>
            </section>
            <script type="application/ld+json">
              {"@type":"Person","name":"Bright Labs"}
            </script>
          </article>
        </body></html>
        """
        people = extract_people_from_html(
            page,
            "https://example.org/content-hub/article",
        )
        self.assertEqual(len(people), 1)
        self.assertEqual(people[0]["person_name"], "Mahesh Perera")
        self.assertEqual(people[0]["linkedin_profile_url"], "")

    def test_writes_joined_people_sheet(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            destination = Path(tmp) / "result.xlsx"
            workbook = Workbook()
            how = workbook.active
            how.title = "How to use"
            how.append(["Guide"])
            organisations = workbook.create_sheet("National non-profits")
            organisations.append(
                [
                    "category",
                    "business_name",
                    "legal_name",
                    "abn",
                    "website",
                    "linkedin_company_url",
                ]
            )
            organisations.append(
                [
                    "nonprofit",
                    "Example Care",
                    "Example Care Ltd",
                    "123 456 789 01",
                    "https://example.org",
                    "https://www.linkedin.com/company/example-care",
                ]
            )
            workbook.save(source)
            organisation_map = organisations_from_workbook(source)
            results = {
                "12345678901": {
                    "status": "people found",
                    "pages_checked": ["https://example.org/team"],
                    "people": [
                        {
                            "person_name": "Jane Smith",
                            "role": "CEO",
                            "linkedin_profile_url": "https://www.linkedin.com/in/jane-smith",
                            "source_type": "official website LinkedIn link",
                            "source_url": "https://example.org/team",
                            "confidence": "high",
                        },
                        {
                            "person_name": "Support Coordination",
                            "role": "Manager",
                            "linkedin_profile_url": "",
                            "source_type": "official website staff page",
                            "source_url": "https://example.org/services",
                            "confidence": "medium",
                        },
                        {
                            "person_name": "Richard Drevid",
                            "role": "Alextina Javi Manager",
                            "linkedin_profile_url": "",
                            "source_type": "official website staff page",
                            "source_url": "https://example.org/team",
                            "confidence": "medium",
                        }
                    ],
                }
            }
            summary = write_enriched_workbook(
                source,
                destination,
                organisation_map,
                results,
            )
            output = load_workbook(destination)
            self.assertIn("Organizations", output.sheetnames)
            self.assertIn("People", output.sheetnames)
            self.assertEqual(output["People"]["D2"].value, "Jane Smith")
            self.assertEqual(output["People"].max_row, 2)
            self.assertEqual(
                output["People"]["F2"].hyperlink.target,
                "https://www.linkedin.com/in/jane-smith",
            )
            self.assertEqual(summary["public_people"], 1)

    def test_merges_verified_public_search_result(self):
        results = {
            "12345678901": {
                "status": "no public staff found",
                "pages_checked": [],
                "people": [],
            }
        }
        merge_public_search_people(
            results,
            [
                {
                    "abn": "123 456 789 01",
                    "person_name": "Jane Smith",
                    "role": "Chief Executive Officer",
                    "linkedin_profile_url": "https://au.linkedin.com/in/jane-smith",
                }
            ],
        )
        person = results["12345678901"]["people"][0]
        self.assertEqual(
            person["linkedin_profile_url"],
            "https://www.linkedin.com/in/jane-smith",
        )
        self.assertEqual(person["source_type"], "publicly indexed LinkedIn result")


if __name__ == "__main__":
    unittest.main()
