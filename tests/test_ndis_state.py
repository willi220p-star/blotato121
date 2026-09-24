import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from lib.ndis_state import parse_location, state_from_postcode, write_state_people_workbook


class NdisStateTest(unittest.TestCase):
    def test_parses_register_head_office(self):
        parsed = parse_location("LARRAKEYAH, NT, 0820, AU")
        self.assertEqual(parsed["state"], "NT")
        self.assertEqual(parsed["state_name"], "Northern Territory")
        self.assertEqual(parsed["suburb"], "LARRAKEYAH")
        self.assertEqual(parsed["postcode"], "0820")

    def test_uses_postcode_when_state_missing(self):
        self.assertEqual(state_from_postcode("2000"), "NSW")
        self.assertEqual(state_from_postcode("3000"), "VIC")
        self.assertEqual(parse_location("MELBOURNE, 3000, AU")["state"], "VIC")

    def test_writes_state_and_people_sheets(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            workbook = Workbook()
            how = workbook.active
            how.title = "How to use"
            how.append(["NDIS providers — disability companies and non-profits"])
            national = workbook.create_sheet("National disability cos")
            national.append(
                [
                    "category",
                    "business_name",
                    "legal_name",
                    "abn",
                    "head_office",
                    "website",
                    "has_nt_presence",
                    "has_darwin_presence",
                ]
            )
            national.append(
                [
                    "disability_company",
                    "2 INCLUDE ENTERPRISE PTY LTD",
                    "2 Include Enterprise Pty Ltd",
                    "50661416468",
                    "PRESTON, VIC, 3072, AU",
                    "https://example.com",
                    "",
                    "",
                ]
            )
            people = workbook.create_sheet("People")
            people.append(
                [
                    "abn",
                    "business_name",
                    "legal_name",
                    "person_name",
                    "role",
                    "linkedin_profile_url",
                    "source_type",
                    "source_url",
                    "confidence",
                ]
            )
            people.append(
                [
                    "50661416468",
                    "2 INCLUDE ENTERPRISE PTY LTD",
                    "2 Include Enterprise Pty Ltd",
                    "Melissa Slimming",
                    "DIRECTOR | BEHAVIOUR SUPPORT PRACTITIONER",
                    None,
                    "official staff page",
                    "https://example.com/team",
                    "high",
                ]
            )
            workbook.save(source)
            destination = Path(tmp) / "out.xlsx"
            summary = write_state_people_workbook(source, destination)
            self.assertEqual(summary["companies"], 1)
            self.assertEqual(summary["people"], 1)
            result = load_workbook(destination)
            national_headers = [cell.value for cell in result["National disability cos"][1]]
            self.assertIn("state", national_headers)
            state_col = national_headers.index("state") + 1
            self.assertEqual(result["National disability cos"].cell(2, state_col).value, "VIC")
            self.assertEqual(result["Company + people"]["A2"].value, "VIC")
            self.assertEqual(result["Company + people"]["H2"].value, "Melissa Slimming")
            self.assertEqual(result["Companies by state"]["H2"].value, 1)
            self.assertIn("Melissa Slimming", result["Companies by state"]["I2"].value)
            self.assertEqual(result["State summary"]["A2"].value, "VIC")
            self.assertEqual(result["State summary"]["C2"].value, 1)


if __name__ == "__main__":
    unittest.main()
