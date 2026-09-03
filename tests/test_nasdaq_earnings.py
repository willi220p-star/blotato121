import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from lib.nasdaq_earnings import (
    COLUMNS,
    bollinger_position,
    build_rows,
    eps_expectation,
    parse_number,
    rating_label,
    scenario_range,
    technical_bias,
    trend_structure,
    write_workbook,
)


class NasdaqEarningsTest(unittest.TestCase):
    def test_parse_number(self):
        self.assertEqual(parse_number("$1,234.50"), 1234.5)
        self.assertEqual(parse_number("($0.04)"), -0.04)
        self.assertIsNone(parse_number("N/A"))

    def test_rating_boundaries(self):
        self.assertEqual(rating_label(0.6), "Strong Buy")
        self.assertEqual(rating_label(0.2), "Buy")
        self.assertEqual(rating_label(0), "Neutral")
        self.assertEqual(rating_label(-0.2), "Sell")
        self.assertEqual(rating_label(-0.7), "Strong Sell")

    def test_trend_and_bollinger(self):
        self.assertEqual(trend_structure(120, 115, 110, 100), "Bullish alignment")
        self.assertEqual(trend_structure(80, 85, 90, 100), "Bearish alignment")
        self.assertEqual(bollinger_position(100, 80, 120), "Middle third")
        self.assertEqual(bollinger_position(125, 80, 120), "At/above upper band")

    def test_scenario_is_symmetric_and_not_negative(self):
        low, high, pct = scenario_range(100, 2)
        self.assertAlmostEqual(100 - low, high - 100)
        self.assertGreater(pct, 0)
        low, _, _ = scenario_range(1, 10)
        self.assertEqual(low, 0)

    def test_eps_expectation(self):
        change, label = eps_expectation(1.2, 1.0)
        self.assertAlmostEqual(change, 20)
        self.assertEqual(label, "Forecast EPS growth")
        _, label = eps_expectation(0.1, -0.2)
        self.assertEqual(label, "Forecast turnaround to profit")

    def test_build_rows_with_technicals(self):
        earnings = [
            {
                "earningsDate": "2026-09-08",
                "time": "time-after-hours",
                "symbol": "DEMO",
                "name": "Demo Corp",
                "marketCap": "$1,000,000",
                "fiscalQuarterEnding": "Aug/2026",
                "epsForecast": "$1.20",
                "noOfEsts": "5",
                "lastYearRptDt": "9/8/2025",
                "lastYearEPS": "$1.00",
            }
        ]
        technicals = {
            "DEMO": {
                "_ticker": "NASDAQ:DEMO",
                "exchange": "NASDAQ",
                "close": 120,
                "change": 1.5,
                "volume": 1000,
                "market_cap_basic": 1000000,
                "price_52_week_high": 130,
                "price_52_week_low": 80,
                "RSI": 60,
                "MACD.macd": 3,
                "MACD.signal": 2,
                "SMA20": 115,
                "SMA50": 110,
                "SMA200": 100,
                "EMA20": 116,
                "EMA50": 111,
                "EMA200": 101,
                "Recommend.All": 0.7,
                "Recommend.MA": 0.8,
                "Recommend.Other": 0.2,
                "ATR": 3,
                "BB.upper": 125,
                "BB.lower": 95,
            }
        }
        row = build_rows(earnings, technicals, "2026-09-03T00:00:00+00:00")[0]
        self.assertEqual(row["report_time"], "After hours")
        self.assertEqual(row["technical_rating"], "Strong Buy")
        self.assertEqual(row["future_bias_not_forecast"], "Bullish technical bias")
        self.assertEqual(row["earnings_expectation"], "Forecast EPS growth")
        self.assertEqual(row["data_status"], "Complete")

    def test_workbook_has_all_and_filtered_sheets(self):
        row = {key: None for key in COLUMNS}
        row.update(
            {
                "earnings_date": "2026-09-08",
                "symbol": "DEMO",
                "company_name": "Demo Corp",
                "report_time": "After hours",
                "consensus_eps_forecast": 1.2,
                "future_bias_not_forecast": "Bullish technical bias",
                "data_status": "Complete",
            }
        )
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "report.xlsx"
            write_workbook(path, [row], date(2026, 9, 4), date(2026, 10, 3))
            wb = load_workbook(path, read_only=True)
            self.assertIn("All upcoming earnings", wb.sheetnames)
            self.assertIn("Forecast profit EPS", wb.sheetnames)
            self.assertIn("Bullish bias", wb.sheetnames)
            self.assertEqual(wb["Forecast profit EPS"].max_row, 2)


if __name__ == "__main__":
    unittest.main()
