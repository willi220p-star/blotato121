import unittest

from lib.arrcs_darwin import (
    contains_person_name,
    darwin_jobs_from_board,
    is_darwin_location,
    parse_job_board_text,
    researched_rows,
    write_outputs,
)


BOARD = """
23 Jobs
Clinical Nurse Manager - Katherine
Full Time
Lead our clinical team in Katherine and support our Aged Care facilities....
Katherine | Northern Territory
Healthcare & Medical, Nursing - Aged Care
First Nations Trainee - Administration
Apprentice/Trainee
An exciting opportunity to join ARRCS as our Administration Trainee.
Darwin | Northern Territory
Administration & Office Support, Administrative Assistants
Executive Manager - Residential Services
Full Time
Lead the future of Aged Care across the Northern Territory.
Darwin | Northern Territory
Healthcare & Medical, Nursing - Aged Care
Caretaker - Alice Springs
Full Time
Join our Alice Springs team.
Alice Springs | Northern Territory
Real Estate & Property, Other
Registered Nurse - Flynn Lodge, Alice Springs
Full Time
A new opportunity to join the team in Flynn Lodge.
Alice Springs | Northern Territory
Healthcare & Medical, Nursing - Aged Care
"""


class ArrcsDarwinTest(unittest.TestCase):
    def test_darwin_location_filter(self):
        self.assertTrue(is_darwin_location("Darwin | Northern Territory"))
        self.assertTrue(is_darwin_location("Coconut Grove, Darwin"))
        self.assertTrue(is_darwin_location("Farrar / Palmerston"))
        self.assertTrue(is_darwin_location("Tiwi, Darwin (Harry's Place)"))
        self.assertFalse(is_darwin_location("Katherine | Northern Territory"))
        self.assertFalse(is_darwin_location("Alice Springs | Northern Territory"))
        self.assertFalse(is_darwin_location("Flynn Lodge, Alice Springs"))
        self.assertFalse(is_darwin_location("Nhulunbuy | Northern Territory"))
        self.assertFalse(is_darwin_location("Pukatja | South Australia"))
        self.assertFalse(is_darwin_location(""))

    def test_job_parser_keeps_darwin_only(self):
        parsed = parse_job_board_text(BOARD)
        titles = [job["title"] for job in parsed]
        self.assertIn("Clinical Nurse Manager - Katherine", titles)
        self.assertIn("First Nations Trainee - Administration", titles)
        darwin = darwin_jobs_from_board(BOARD)
        self.assertEqual(
            [job["title"] for job in darwin],
            ["First Nations Trainee - Administration", "Executive Manager - Residential Services"],
        )
        self.assertTrue(all(is_darwin_location(job["location"]) for job in darwin))

    def test_export_has_no_people_and_covers_timeframes(self):
        rows = researched_rows()
        self.assertGreaterEqual(len(rows), 40)
        timeframes = {row["timeframe"] for row in rows}
        self.assertEqual(timeframes, {"past", "present", "future"})
        live = [row for row in rows if row["status"] == "live vacancy"]
        self.assertEqual(len(live), 2)
        self.assertTrue(any("9792357" in row["source_url"] for row in live))
        self.assertTrue(any("9784188" in row["source_url"] for row in live))
        for row in rows:
            blob = " ".join(str(v) for v in row.values())
            self.assertFalse(contains_person_name(blob), blob)
            if row["timeframe"] in {"present", "past"}:
                self.assertTrue(
                    is_darwin_location(row["location"]) or "Darwin" in row["location"],
                    row["location"],
                )

    def test_excludes_family_support_and_flynn(self):
        blob = " ".join(" ".join(row.values()) for row in researched_rows()).lower()
        self.assertNotIn("mutitjulu itiku", blob)
        self.assertNotIn("school nutrition", blob)
        self.assertNotIn("flynn lodge", blob)

    def test_workbook_roundtrip(self):
        import tempfile
        from pathlib import Path

        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            payload = write_outputs(Path(tmp))
            path = Path(payload["xlsx"])
            self.assertTrue(path.is_file())
            wb = load_workbook(path)
            self.assertEqual(
                wb.sheetnames,
                ["How to use", "Present", "Past", "Future", "All in this list", "Sources and gaps"],
            )
            present = wb["Present"]
            headers = [cell.value for cell in present[1]]
            self.assertEqual(headers, ["timeframe", "team", "position", "location", "status", "source_url", "notes"])
            self.assertGreater(present.max_row, 10)
            csv_text = (Path(tmp) / "arrcs-darwin-teams.csv").read_text(encoding="utf-8")
            self.assertIn("First Nations trainee", csv_text)
            self.assertNotIn("Jenny Messell", csv_text)
            self.assertNotIn("Wendy Hubbard", csv_text)


if __name__ == "__main__":
    unittest.main()
