import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from lib.linkedin_enrichment import (
    enrich_workbook,
    extract_linkedin_urls,
    linkedin_profile_name,
    normalise_linkedin_url,
    nonprofit_organisations,
)


class LinkedInEnrichmentTest(unittest.TestCase):
    def test_extracts_company_pages_and_skips_people(self):
        html = """
        <a href="https://au.linkedin.com/company/autism-nt/?trk=footer">LinkedIn</a>
        <a href="https://www.linkedin.com/in/person-name">Person</a>
        <script>{"url":"https:\\/\\/www.linkedin.com\\/company\\/carers-nt\\/posts"}</script>
        """
        self.assertEqual(
            extract_linkedin_urls(html),
            [
                "https://www.linkedin.com/company/autism-nt",
                "https://www.linkedin.com/company/carers-nt",
            ],
        )

    def test_normalises_only_organisation_profiles(self):
        self.assertEqual(
            normalise_linkedin_url("https://linkedin.com/company/mjd-foundation/?trk=x"),
            "https://www.linkedin.com/company/mjd-foundation",
        )
        self.assertIsNone(normalise_linkedin_url("https://www.linkedin.com/in/a-person"))
        self.assertIsNone(normalise_linkedin_url("https://example.com/company/demo"))
        self.assertIsNone(normalise_linkedin_url("https://www.linkedin.com/company/squarespace"))
        self.assertIsNone(normalise_linkedin_url("https://www.linkedin.com/company/l"))
        self.assertEqual(linkedin_profile_name("https://www.linkedin.com/company/mjd-foundation"), "Mjd Foundation")

    def test_enriches_nonprofit_rows_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            out = Path(tmp) / "out.xlsx"
            wb = Workbook()
            how = wb.active
            how.title = "How to use"
            how["A1"] = "Guide"
            ws = wb.create_sheet("National non-profits")
            ws.append(["category", "business_name", "abn", "website"])
            ws.append(["nonprofit", "Autism NT", "123", "https://autismnt.org.au"])
            ws.append(["disability_company", "Company", "456", "https://company.example"])
            wb.save(source)

            organisations = nonprofit_organisations(source)
            self.assertEqual(organisations, {"123": "https://autismnt.org.au"})
            summary = enrich_workbook(
                source,
                out,
                {
                    "123": {
                        "url": "https://www.linkedin.com/company/autism-nt",
                        "profile_name": "Autism Nt",
                        "source": "https://autismnt.org.au",
                        "status": "confirmed: published on official website",
                    }
                },
            )
            self.assertEqual(summary["confirmed_links"], 1)
            result = load_workbook(out)
            sheet = result["National non-profits"]
            headers = [cell.value for cell in sheet[1]]
            link_col = headers.index("linkedin_company_url") + 1
            status_col = headers.index("linkedin_status") + 1
            self.assertEqual(sheet.cell(2, link_col).value, "https://www.linkedin.com/company/autism-nt")
            self.assertEqual(sheet.cell(2, status_col).value, "confirmed: published on official website")
            self.assertIsNone(sheet.cell(3, link_col).value)


if __name__ == "__main__":
    unittest.main()
