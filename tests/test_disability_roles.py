import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from lib.disability_roles import (
    classify_disability_role,
    extract_job_openings,
    plausible_job_title,
    write_role_workbook,
)


class DisabilityRolesTest(unittest.TestCase):
    def test_classifies_requested_disability_roles(self):
        cases = {
            "Disability Support Worker": "Disability Support Worker",
            "Case Manager": "Case Manager / Case Worker",
            "Senior Occupational Therapist": "Occupational Therapist",
            "Positive Behaviour Support Practitioner": "Behaviour Support Practitioner",
            "Psychosocial Recovery Coach": "Psychosocial Recovery Coach",
            "Support Coordination Team Leader": "Support Coordinator",
            "Community Team Leader": "Team Leader",
            "Care Coordinator": "Disability / Service Coordinator",
        }
        for title, expected in cases.items():
            with self.subTest(title=title):
                self.assertEqual(classify_disability_role(title), expected)

    def test_excludes_corporate_coordinators_and_leaders(self):
        self.assertEqual(classify_disability_role("Finance Coordinator"), "")
        self.assertEqual(classify_disability_role("Marketing Team Leader"), "")
        self.assertEqual(classify_disability_role("Project Coordinator"), "")

    def test_rejects_services_stories_and_questions_as_openings(self):
        rejected = (
            "Support Coordination",
            "Occupational Therapy",
            "Positive Behaviour Support",
            "The Essential Role of Male Disability Support Workers",
            "Do I need experience in disability support to apply?",
            "A rewarding career: Mukti's story as a Home Support Worker",
            "Allied Health Jobs",
        )
        for title in rejected:
            with self.subTest(title=title):
                self.assertFalse(plausible_job_title(title))
        self.assertTrue(plausible_job_title("Disability Support Worker"))
        self.assertTrue(plausible_job_title("Senior Occupational Therapist"))

    def test_extracts_current_disability_job_postings(self):
        page = """
        <html><body>
          <script type="application/ld+json">
          [
            {
              "@type":"JobPosting",
              "title":"Disability Support Worker",
              "url":"/jobs/support-worker",
              "datePosted":"2026-09-01",
              "validThrough":"2026-12-31",
              "employmentType":"PART_TIME",
              "jobLocation":{"address":{
                "addressLocality":"Darwin",
                "addressRegion":"NT",
                "addressCountry":"AU"
              }}
            },
            {
              "@type":"JobPosting",
              "title":"Case Manager",
              "url":"/jobs/old",
              "validThrough":"2020-01-01"
            },
            {
              "@type":"JobPosting",
              "title":"Accountant",
              "url":"/jobs/accountant"
            }
          ]
          </script>
          <a href="/careers/occupational-therapist">Occupational Therapist</a>
          <a href="/careers/occupational-therapist">Occupational therapy jobs</a>
        </body></html>
        """
        openings = extract_job_openings(page, "https://example.org/careers")
        titles = {opening["position_title"] for opening in openings}
        self.assertEqual(
            titles,
            {"Disability Support Worker", "Occupational Therapist"},
        )
        support = next(
            opening
            for opening in openings
            if opening["position_title"] == "Disability Support Worker"
        )
        self.assertEqual(support["location"], "Darwin, NT, AU")

    def test_does_not_treat_service_links_as_job_openings(self):
        page = """
        <html><body>
          <a href="/services/support-coordination">Support Coordination</a>
          <a href="/services/occupational-therapy">Occupational Therapy</a>
        </body></html>
        """
        self.assertEqual(
            extract_job_openings(page, "https://example.org/"),
            [],
        )
        self.assertEqual(
            extract_job_openings(page, "https://example.org/careers"),
            [],
        )

    def test_writes_three_role_tabs_and_summary_counts(self):
        with TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            destination = Path(tmp) / "result.xlsx"
            workbook = Workbook()
            workbook.active.title = "How to use"
            workbook["How to use"].append(["Guide"])
            workbook.save(source)
            organisations = {
                "123": {
                    "abn": "123",
                    "business_name": "Example Care",
                    "legal_name": "Example Care Ltd",
                }
            }
            people = [
                {
                    "role_category": "Disability Support Worker",
                    "abn": "123",
                    "business_name": "Example Care",
                    "legal_name": "Example Care Ltd",
                    "person_name": "Jane Smith",
                    "role": "Disability Support Worker",
                    "linkedin_profile_url": "https://www.linkedin.com/in/jane-smith",
                    "source_type": "publicly indexed LinkedIn result",
                    "source_url": "https://www.linkedin.com/in/jane-smith",
                    "confidence": "high",
                }
            ]
            jobs = {
                "123": {
                    "status": "positions found",
                    "openings": [
                        {
                            "role_category": "Occupational Therapist",
                            "position_title": "Occupational Therapist",
                            "status": "Published JobPosting",
                            "date_posted": "2026-09-01",
                            "valid_through": "2026-12-31",
                            "location": "Darwin, NT, AU",
                            "employment_type": "FULL_TIME",
                            "job_url": "https://example.org/jobs/ot",
                            "source_page": "https://example.org/careers",
                        }
                    ],
                }
            }
            summary = write_role_workbook(
                source,
                destination,
                organisations,
                people,
                jobs,
            )
            output = load_workbook(destination)
            self.assertIn("Disability Role People", output.sheetnames)
            self.assertIn("Disability Job Openings", output.sheetnames)
            self.assertIn("Company Role Summary", output.sheetnames)
            self.assertEqual(output["Company Role Summary"]["G2"].value, 1)
            self.assertEqual(output["Company Role Summary"]["H2"].value, 1)
            self.assertEqual(output["Company Role Summary"]["I2"].value, 2)
            self.assertEqual(summary["published_openings"], 1)


if __name__ == "__main__":
    unittest.main()
