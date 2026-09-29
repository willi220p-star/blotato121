"""CSV, XLSX, and markdown market summary."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.darwin_it.models import COMPANY_COLUMNS, SOURCE_COLUMNS, UNKNOWN, VACANCY_COLUMNS, today_iso


def completeness(companies: list[dict], vacancies: list[dict]) -> dict:
    def has(row: dict, field: str) -> bool:
        value = str(row.get(field) or "")
        return value not in {"", UNKNOWN, "Not found", "NO"}

    company_vacs = [row for row in vacancies if row.get("vacancy_source") == "Company Career Page"]
    seek_vacs = [row for row in vacancies if row.get("vacancy_source") == "SEEK"]
    both = 0
    for company in companies:
        if int(company.get("company_vacancy_count") or 0) and int(company.get("seek_vacancy_count") or 0):
            both += 1
    return {
        "total_companies_found": len(companies),
        "verified_companies": sum(1 for row in companies if row.get("research_confidence") in {"High", "Medium"}),
        "companies_with_websites": sum(1 for row in companies if has(row, "company_website")),
        "companies_with_linkedin": sum(1 for row in companies if has(row, "linkedin_url")),
        "companies_with_employee_data": sum(1 for row in companies if has(row, "employee_count")),
        "companies_with_career_pages": sum(1 for row in companies if row.get("career_page_found") == "YES"),
        "companies_with_active_vacancies": sum(1 for row in companies if int(row.get("total_vacancy_count") or 0) > 0),
        "companies_with_SEEK_vacancies": sum(1 for row in companies if int(row.get("seek_vacancy_count") or 0) > 0),
        "companies_with_both_company_and_SEEK_vacancies": both,
        "companies_without_verified_employee_count": sum(1 for row in companies if not has(row, "employee_count")),
        "companies_without_career_page": sum(1 for row in companies if row.get("career_page_found") != "YES"),
        "official_career_vacancies": len(company_vacs),
        "seek_vacancies": len(seek_vacs),
        "research_date": today_iso(),
    }


def _md_cell(value) -> str:
    return str(value or "").replace("|", "/").replace("\n", " ")


def summary_markdown(companies: list[dict], vacancies: list[dict], stats: dict) -> str:
    lines = [
        "# Darwin IT market summary",
        "",
        f"Research date: {stats['research_date']}",
        "",
        "Public-source research of IT, ICT, technology, software, cybersecurity, MSP and related companies with a Darwin / Greater Darwin presence. Employee counts and vacancies are recorded only when a public source exists. Unknown values are marked, not guessed. SEEK individual job pages were not fetched. LinkedIn was not crawled.",
        "",
        "## Darwin IT Market",
        "",
        f"Total companies discovered: **{stats['total_companies_found']}**",
        "",
        f"Verified companies: **{stats['verified_companies']}**",
        "",
        f"Companies with employee data: **{stats['companies_with_employee_data']}**",
        "",
        f"Companies with active vacancies: **{stats['companies_with_active_vacancies']}**",
        "",
        f"Companies with SEEK vacancies: **{stats['companies_with_SEEK_vacancies']}**",
        "",
        f"Companies with official career vacancies: **{sum(1 for row in companies if int(row.get('company_vacancy_count') or 0) > 0)}**",
        "",
        "## Completeness",
        "",
        f"- companies_with_websites: {stats['companies_with_websites']}",
        f"- companies_with_linkedin: {stats['companies_with_linkedin']}",
        f"- companies_with_career_pages: {stats['companies_with_career_pages']}",
        f"- companies_with_both_company_and_SEEK_vacancies: {stats['companies_with_both_company_and_SEEK_vacancies']}",
        f"- companies_without_verified_employee_count: {stats['companies_without_verified_employee_count']}",
        f"- companies_without_career_page: {stats['companies_without_career_page']}",
        "",
        "## Company Table",
        "",
        "| Company | Description | Employees | Website | LinkedIn | Careers | Vacancies | SEEK |",
        "| ------- | ----------- | --------: | ------- | -------- | ------- | --------: | ---- |",
    ]
    for row in companies:
        lines.append(
            "| "
            + " | ".join(
                [
                    _md_cell(row.get("company_name")),
                    _md_cell((row.get("company_description") or "")[:140]),
                    _md_cell(row.get("employee_count")),
                    _md_cell(row.get("company_website")),
                    _md_cell(row.get("linkedin_url")),
                    _md_cell(row.get("career_page_found")),
                    _md_cell(row.get("total_vacancy_count")),
                    _md_cell(row.get("seek_vacancy_count")),
                ]
            )
            + " |"
        )
    lines += [
        "",
        "## Vacancy Table",
        "",
        "| Company | Job | Location | Type | Source | URL |",
        "| ------- | --- | -------- | ---- | ------ | --- |",
    ]
    if vacancies:
        for row in vacancies:
            lines.append(
                "| "
                + " | ".join(
                    [
                        _md_cell(row.get("company_name")),
                        _md_cell(row.get("vacancy_title")),
                        _md_cell(row.get("vacancy_location")),
                        _md_cell(row.get("employment_type")),
                        _md_cell(row.get("vacancy_source")),
                        _md_cell(row.get("vacancy_url")),
                    ]
                )
                + " |"
            )
    else:
        lines.append("| None verified | — | — | — | — | — |")
    lines += [
        "",
        "## Method",
        "",
        "1. Discover companies from ICTNT, Chamber of Commerce NT, InfoMSP, local directories, and official-website search.",
        "2. Deduplicate by official domain, name, LinkedIn slug, and ABN where present.",
        "3. Verify each official website (directory URLs are not stored as the company website).",
        "4. Record LinkedIn only when published on the official site or as a public search snippet with a source URL.",
        "5. Record employee counts only from sourced evidence. Ranges stay ranges.",
        "6. Check official career pages. Extract JobPosting JSON-LD and career-page job links only.",
        "7. Attempt SEEK keyword search only. Do not fetch `/job/` listings or bypass blocks.",
        "",
    ]
    return "\n".join(lines) + "\n"


def _write_csv(path: Path, rows: list[dict], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in columns})


def _style_sheet(sheet, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(1, index, header)
        cell.fill = fill
        cell.font = font
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 16), 48)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_xlsx(path: Path, companies: list[dict], vacancies: list[dict], sources: list[dict], stats: dict) -> None:
    workbook = Workbook()
    how = workbook.active
    how.title = "How to use"
    how.append(["Darwin IT company and hiring database"])
    how.append([f"Research date: {stats['research_date']}"])
    how.append(["Companies are included only with evidence of Darwin / Greater Darwin presence or Darwin IT vacancies."])
    how.append(["Unknown means the field could not be verified. Values are never invented."])
    how.append(["SEEK: keyword search only. Individual SEEK job pages are not fetched. LinkedIn is not crawled."])
    how.append(["Sheets: Companies, Vacancies, Sources, Completeness."])
    sheets = (
        ("Companies", companies, COMPANY_COLUMNS),
        ("Vacancies", vacancies, VACANCY_COLUMNS),
        ("Sources", sources, SOURCE_COLUMNS),
        ("Completeness", [{"metric": key, "value": value} for key, value in stats.items()], ["metric", "value"]),
    )
    for title, rows, columns in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(columns)
        for row in rows:
            sheet.append([row.get(column, "") for column in columns])
        _style_sheet(sheet, columns)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_outputs(
    companies: list[dict],
    vacancies: list[dict],
    sources: list[dict],
    *,
    data_dir: Path,
    reports_dir: Path,
    feeds_dir: Path,
) -> dict:
    stats = completeness(companies, vacancies)
    data_dir.mkdir(parents=True, exist_ok=True)
    reports_dir.mkdir(parents=True, exist_ok=True)
    feeds_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(data_dir / "companies.csv", companies, COMPANY_COLUMNS)
    _write_csv(data_dir / "vacancies.csv", vacancies, VACANCY_COLUMNS)
    _write_csv(data_dir / "sources.csv", sources, SOURCE_COLUMNS)
    _write_csv(feeds_dir / "darwin-it-companies.csv", companies, COMPANY_COLUMNS)
    _write_csv(feeds_dir / "darwin-it-vacancies.csv", vacancies, VACANCY_COLUMNS)
    _write_csv(feeds_dir / "darwin-it-sources.csv", sources, SOURCE_COLUMNS)
    (feeds_dir / "darwin-it-market-summary.md").write_text(
        summary_markdown(companies, vacancies, stats),
        encoding="utf-8",
    )
    (reports_dir / "darwin-it-market-summary.md").write_text(
        summary_markdown(companies, vacancies, stats),
        encoding="utf-8",
    )
    write_xlsx(feeds_dir / "Darwin_IT_company_database.xlsx", companies, vacancies, sources, stats)
    write_xlsx(data_dir / "Darwin_IT_company_database.xlsx", companies, vacancies, sources, stats)
    return stats
