"""CSV, JSON, SQLite, Excel and markdown summary exports."""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.nt_ndis.config import UNKNOWN, today_iso
from src.nt_ndis.database import write_sqlite

COMPANY_FIELDS = [
    "company_id",
    "company_name",
    "trading_name",
    "legal_name",
    "abn",
    "org_type",
    "description",
    "services",
    "locations",
    "address",
    "phone",
    "website",
    "email",
    "linkedin",
    "employee_count",
    "employee_count_min",
    "employee_count_max",
    "employee_count_source",
    "employee_count_source_url",
    "ndis_provider",
    "ndis_evidence",
    "nt_operation_verified",
    "scrape_status",
    "research_confidence",
    "source_urls",
    "last_verified",
]
EMPLOYEE_FIELDS = [
    "company_id",
    "company_name",
    "abn",
    "first_name",
    "last_name",
    "full_name",
    "job_title",
    "department",
    "location",
    "linkedin_url",
    "work_email",
    "email_type",
    "email_verified",
    "email_status",
    "source_url",
    "source_type",
    "verification_status",
    "priority_role",
]
VACANCY_FIELDS = [
    "company_id",
    "company_name",
    "job_title",
    "location",
    "employment_type",
    "description",
    "date_posted",
    "job_url",
    "source",
    "active",
    "last_checked",
]
SOURCE_FIELDS = ["company_id", "employee_id", "vacancy_id", "source_type", "source_url", "source_title", "retrieved_at"]
COMBINED_FIELDS = [
    "company_name",
    "abn",
    "org_type",
    "locations",
    "employee_count",
    "employee_name",
    "job_title",
    "email",
    "website",
    "linkedin",
    "phone",
    "vacancy",
    "vacancy_url",
    "source",
]


def _write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: row.get(field, "") for field in fields})


def combined_rows(companies: list[dict], employees: list[dict], vacancies: list[dict]) -> list[dict]:
    by_company: dict[str, list[dict]] = {}
    for person in employees:
        by_company.setdefault(str(person.get("company_id")), []).append(person)
    vacs: dict[str, list[dict]] = {}
    for job in vacancies:
        vacs.setdefault(str(job.get("company_id")), []).append(job)
    rows = []
    for company in companies:
        cid = str(company.get("company_id"))
        people = by_company.get(cid) or [None]
        jobs = vacs.get(cid) or [None]
        for person in people:
            for job in jobs[:1] if person is None else [jobs[0] if jobs and jobs[0] else None]:
                rows.append(
                    {
                        "company_name": company.get("company_name"),
                        "abn": company.get("abn"),
                        "org_type": company.get("org_type"),
                        "locations": company.get("locations"),
                        "employee_count": company.get("employee_count"),
                        "employee_name": (person or {}).get("full_name") or "",
                        "job_title": (person or {}).get("job_title") or "",
                        "email": (person or {}).get("work_email") or company.get("email") or "",
                        "website": company.get("website"),
                        "linkedin": (person or {}).get("linkedin_url") or company.get("linkedin"),
                        "phone": company.get("phone"),
                        "vacancy": (job or {}).get("job_title") or "",
                        "vacancy_url": (job or {}).get("job_url") or "",
                        "source": company.get("source_urls"),
                    }
                )
    return rows


def completeness(companies: list[dict], employees: list[dict], vacancies: list[dict]) -> dict:
    def has(row: dict, field: str) -> bool:
        return str(row.get(field) or "") not in {"", UNKNOWN, "Unknown", "Not found"}

    loc_counts: Counter[str] = Counter()
    for company in companies:
        for part in str(company.get("locations") or "").split(","):
            label = part.strip()
            if label:
                loc_counts[label] += 1
    service_counts: Counter[str] = Counter()
    for company in companies:
        for part in str(company.get("services") or "").split(";"):
            label = part.strip()
            if label and label != UNKNOWN:
                service_counts[label] += 1
    return {
        "research_date": today_iso(),
        "total_companies_discovered": len(companies),
        "total_verified_nt_ndis_providers": sum(1 for row in companies if row.get("nt_operation_verified") == "yes"),
        "private_companies": sum(1 for row in companies if row.get("org_type") == "private_company"),
        "nonprofit_ndis_providers": sum(1 for row in companies if row.get("org_type") == "nonprofit_ndis"),
        "sole_traders": sum(1 for row in companies if row.get("org_type") == "sole_trader"),
        "total_employees_found": len(employees),
        "total_public_work_emails": sum(1 for row in companies if has(row, "email"))
        + sum(1 for row in employees if has(row, "work_email")),
        "total_companies_with_vacancies": len({row.get("company_id") for row in vacancies}),
        "total_active_vacancies": len(vacancies),
        "total_companies_without_public_employee_information": sum(
            1
            for row in companies
            if not any(person.get("company_id") == row.get("company_id") for person in employees)
        ),
        "companies_with_websites": sum(1 for row in companies if has(row, "website")),
        "companies_with_employee_data": sum(1 for row in companies if has(row, "employee_count")),
        "top_locations": loc_counts.most_common(10),
        "top_services": service_counts.most_common(12),
        "vacancies_by_title": Counter(row.get("job_title") or UNKNOWN for row in vacancies).most_common(12),
        "employee_size_ranges": Counter(
            "Unknown"
            if not has(row, "employee_count")
            else (
                "1-10"
                if str(row.get("employee_count_max") or row.get("employee_count")).split("-")[-1].isdigit()
                and int(str(row.get("employee_count_max") or row.get("employee_count")).split("-")[-1]) <= 10
                else "11+"
            )
            for row in companies
        ),
    }


def summary_markdown(stats: dict, companies: list[dict], vacancies: list[dict]) -> str:
    lines = [
        "# NT private NDIS provider intelligence",
        "",
        f"Research date: {stats['research_date']}",
        "",
        "Approved NDIS providers with Northern Territory outlets, from the official NDIS Commission register plus official-website enrichment. Government departments, hospitals and councils are excluded. Employee names, emails and vacancies are recorded only when published on an official company page. LinkedIn is not crawled. SEEK `/job/` pages are not fetched.",
        "",
        f"Total companies discovered: **{stats['total_companies_discovered']}**",
        "",
        f"Total verified NT NDIS providers: **{stats['total_verified_nt_ndis_providers']}**",
        "",
        f"Private companies (Pty Ltd): **{stats['private_companies']}**",
        "",
        f"Total employees found: **{stats['total_employees_found']}**",
        "",
        f"Total public work emails: **{stats['total_public_work_emails']}**",
        "",
        f"Total companies with vacancies: **{stats['total_companies_with_vacancies']}**",
        "",
        f"Total active vacancies: **{stats['total_active_vacancies']}**",
        "",
        f"Total companies without public employee information: **{stats['total_companies_without_public_employee_information']}**",
        "",
        "## Top NT locations by providers",
        "",
    ]
    for label, count in stats.get("top_locations") or []:
        lines.append(f"- {label}: {count}")
    lines += ["", "## Providers by service type", ""]
    for label, count in stats.get("top_services") or []:
        lines.append(f"- {label}: {count}")
    lines += ["", "## Vacancies by job title", ""]
    if stats.get("vacancies_by_title"):
        for label, count in stats["vacancies_by_title"]:
            lines.append(f"- {label}: {count}")
    else:
        lines.append("- None verified on official career pages")
    lines += ["", "## Sample companies", "", "| Company | Type | Locations | Website | Employees |", "| --- | --- | --- | --- | ---: |"]
    for row in companies[:25]:
        lines.append(
            f"| {row.get('company_name')} | {row.get('org_type')} | {row.get('locations')} | {row.get('website')} | {row.get('employee_count')} |"
        )
    return "\n".join(lines) + "\n"


def _xlsx(path: Path, companies: list[dict], employees: list[dict], vacancies: list[dict], sources: list[dict], combined: list[dict]) -> None:
    workbook = Workbook()
    how = workbook.active
    how.title = "How to use"
    how.append(["NT private NDIS provider intelligence"])
    how.append(["Official NDIS Commission register + official websites. No LinkedIn crawl. No invented emails."])
    sheets = (
        ("Companies", companies, COMPANY_FIELDS),
        ("Employees", employees, EMPLOYEE_FIELDS),
        ("Vacancies", vacancies, VACANCY_FIELDS),
        ("Sources", sources, SOURCE_FIELDS),
        ("Combined", combined, COMBINED_FIELDS),
    )
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for title, rows, fields in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(fields)
        for row in rows:
            sheet.append([row.get(field, "") for field in fields])
        for index, header in enumerate(fields, start=1):
            cell = sheet.cell(1, index)
            cell.fill = fill
            cell.font = font
            sheet.column_dimensions[get_column_letter(index)].width = min(max(len(header) + 2, 14), 42)
        sheet.auto_filter.ref = sheet.dimensions
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_outputs(
    companies: list[dict],
    employees: list[dict],
    vacancies: list[dict],
    sources: list[dict],
    *,
    output_dir: Path,
    data_dir: Path,
    feeds_dir: Path,
    reports_dir: Path,
) -> dict:
    stats = completeness(companies, employees, vacancies)
    combined = combined_rows(companies, employees, vacancies)
    for folder in (output_dir, data_dir, feeds_dir, reports_dir):
        folder.mkdir(parents=True, exist_ok=True)
    _write_csv(output_dir / "companies.csv", companies, COMPANY_FIELDS)
    _write_csv(output_dir / "employees.csv", employees, EMPLOYEE_FIELDS)
    _write_csv(output_dir / "vacancies.csv", vacancies, VACANCY_FIELDS)
    _write_csv(output_dir / "sources.csv", sources, SOURCE_FIELDS)
    _write_csv(output_dir / "combined.csv", combined, COMBINED_FIELDS)
    _write_csv(data_dir / "companies.csv", companies, COMPANY_FIELDS)
    _write_csv(data_dir / "employees.csv", employees, EMPLOYEE_FIELDS)
    _write_csv(data_dir / "vacancies.csv", vacancies, VACANCY_FIELDS)
    _write_csv(data_dir / "sources.csv", sources, SOURCE_FIELDS)
    _write_csv(feeds_dir / "nt-ndis-private-companies.csv", companies, COMPANY_FIELDS)
    _write_csv(feeds_dir / "nt-ndis-private-employees.csv", employees, EMPLOYEE_FIELDS)
    _write_csv(feeds_dir / "nt-ndis-private-vacancies.csv", vacancies, VACANCY_FIELDS)
    _write_csv(feeds_dir / "nt-ndis-private-combined.csv", combined, COMBINED_FIELDS)
    payload = {"stats": stats, "companies": companies, "employees": employees, "vacancies": vacancies, "sources": sources}
    (output_dir / "nt-ndis-intelligence.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    (feeds_dir / "nt-ndis-intelligence.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    write_sqlite(output_dir / "nt-ndis-intelligence.sqlite", companies, employees, vacancies, sources)
    write_sqlite(data_dir / "nt-ndis-intelligence.sqlite", companies, employees, vacancies, sources)
    markdown = summary_markdown(stats, companies, vacancies)
    (reports_dir / "nt-ndis-private-summary.md").write_text(markdown, encoding="utf-8")
    (feeds_dir / "nt-ndis-private-summary.md").write_text(markdown, encoding="utf-8")
    _xlsx(feeds_dir / "NT_private_NDIS_provider_intelligence.xlsx", companies, employees, vacancies, sources, combined)
    _xlsx(output_dir / "NT_private_NDIS_provider_intelligence.xlsx", companies, employees, vacancies, sources, combined)
    return stats
