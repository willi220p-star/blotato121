"""Discover organisation LinkedIn pages from their official websites.

LinkedIn itself is never fetched. A URL is accepted only when it is published
by the organisation's website (or already supplied as its official website).
"""

from __future__ import annotations

import html
import re
import threading
import urllib.robotparser
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urljoin, urlparse, urlunparse

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

USER_AGENT = "Mozilla/5.0 (compatible; DGK-NDIS-LinkedIn-discovery/1.0)"
LINKEDIN_COLUMNS = (
    "linkedin_company_url",
    "linkedin_profile_name",
    "linkedin_source",
    "linkedin_status",
)
LINKEDIN_RE = re.compile(
    r"""https?://(?:[a-z]{2,3}\.)?linkedin\.com/
        (?P<kind>company|showcase|school)/
        (?P<slug>[A-Za-z0-9_%.\-]+)
        (?:/[^\s"'<>\\]*)?""",
    re.I | re.X,
)
SOCIAL_REDIRECT_RE = re.compile(
    r"""(?:href|url|target|redirect|destination)=
        (?P<url>https?%3A%2F%2F(?:www\.)?linkedin\.com%2F(?:company|showcase|school)%2F[^&"'<>]+)""",
    re.I | re.X,
)
NON_ORGANISATION_SLUGS = {
    "wix-com",
    "squarespace",
    "wordpress",
    "wordpress-com",
    "webflow",
    "shopify",
    "yourpage",
}
_local = threading.local()


def _session() -> requests.Session:
    session = getattr(_local, "session", None)
    if session is None:
        session = requests.Session()
        session.headers.update({"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml"})
        _local.session = session
    return session


def normalise_linkedin_url(url: str) -> str | None:
    value = html.unescape((url or "").replace("\\/", "/").strip())
    if not value:
        return None
    if value.startswith("//"):
        value = "https:" + value
    parsed = urlparse(value)
    host = parsed.netloc.lower().split(":", 1)[0]
    if host not in {"linkedin.com", "www.linkedin.com"} and not host.endswith(".linkedin.com"):
        return None
    match = re.match(r"^/(company|showcase|school)/([^/?#]+)", parsed.path, re.I)
    if not match:
        return None
    kind, slug = match.groups()
    if (
        not slug
        or slug.lower() in {"login", "signup", "undefined"}
        or slug.lower() in NON_ORGANISATION_SLUGS
        or (not slug.isdigit() and len(slug) < 3)
    ):
        return None
    return f"https://www.linkedin.com/{kind.lower()}/{slug.strip('/')}"


def linkedin_profile_name(url: str | None) -> str:
    if not url:
        return ""
    slug = urlparse(url).path.rstrip("/").split("/")[-1]
    return re.sub(r"[-_]+", " ", slug).strip().title()


def extract_linkedin_urls(content: str) -> list[str]:
    decoded = html.unescape((content or "").replace("\\/", "/"))
    found: list[str] = []
    for match in LINKEDIN_RE.finditer(decoded):
        candidate = normalise_linkedin_url(match.group(0))
        if candidate and candidate not in found:
            found.append(candidate)
    # Some sites wrap social links in a URL-encoded redirect.
    from urllib.parse import unquote

    for match in SOCIAL_REDIRECT_RE.finditer(decoded):
        candidate = normalise_linkedin_url(unquote(match.group("url")))
        if candidate and candidate not in found:
            found.append(candidate)
    return found


def _robots_allows(url: str, timeout: int) -> tuple[bool, str]:
    parsed = urlparse(url)
    robots_url = f"{parsed.scheme}://{parsed.netloc}/robots.txt"
    try:
        response = _session().get(robots_url, timeout=timeout, allow_redirects=True)
        if response.status_code in {401, 403}:
            return False, f"robots HTTP {response.status_code}"
        if response.status_code >= 400:
            return True, f"robots unavailable HTTP {response.status_code}"
        parser = urllib.robotparser.RobotFileParser()
        parser.set_url(robots_url)
        parser.parse(response.text.splitlines())
        return parser.can_fetch(USER_AGENT, url), "robots checked"
    except requests.RequestException as exc:
        return True, f"robots unavailable: {type(exc).__name__}"


def _candidate_pages(home_url: str, home_html: str) -> list[str]:
    """Return a few same-host pages likely to contain footer/social links."""
    parsed = urlparse(home_url)
    origin = f"{parsed.scheme}://{parsed.netloc}"
    candidates = [home_url]
    for href in re.findall(r"""href\s*=\s*["']([^"'#]+)""", home_html or "", re.I):
        target = urljoin(home_url, html.unescape(href))
        target_parsed = urlparse(target)
        if target_parsed.netloc.lower() != parsed.netloc.lower():
            continue
        low = target_parsed.path.lower().rstrip("/")
        if any(hint in low for hint in ("/about", "/contact", "/connect", "/social")):
            clean = urlunparse((target_parsed.scheme, target_parsed.netloc, target_parsed.path, "", "", ""))
            if clean not in candidates:
                candidates.append(clean)
        if len(candidates) >= 3:
            break
    return candidates


def discover_linkedin(website: str, timeout: int = 10) -> dict:
    website = (website or "").strip()
    if not website:
        return {"url": "", "profile_name": "", "source": "", "status": "no official website"}
    if "linkedin.com" in website.lower():
        url = normalise_linkedin_url(website)
        return {
            "url": url or "",
            "profile_name": linkedin_profile_name(url),
            "source": website if url else "",
            "status": "confirmed: NDIS register website" if url else "website is not an organisation LinkedIn page",
        }
    if not website.startswith(("http://", "https://")):
        website = "https://" + website
    allowed, robots_note = _robots_allows(website, timeout)
    if not allowed:
        return {"url": "", "profile_name": "", "source": website, "status": f"blocked by robots.txt ({robots_note})"}
    try:
        response = _session().get(website, timeout=timeout, allow_redirects=True)
        if response.status_code >= 400:
            return {
                "url": "",
                "profile_name": "",
                "source": response.url or website,
                "status": f"official website HTTP {response.status_code}",
            }
        content_type = response.headers.get("Content-Type", "")
        if "html" not in content_type.lower() and "<html" not in response.text[:1000].lower():
            return {"url": "", "profile_name": "", "source": response.url, "status": "official website is not HTML"}
        pages = _candidate_pages(response.url, response.text)
        for index, page_url in enumerate(pages):
            page_text = response.text if index == 0 else ""
            if index:
                page_allowed, _ = _robots_allows(page_url, timeout)
                if not page_allowed:
                    continue
                try:
                    page_response = _session().get(page_url, timeout=timeout, allow_redirects=True)
                    if page_response.status_code >= 400:
                        continue
                    page_text = page_response.text
                    page_url = page_response.url
                except requests.RequestException:
                    continue
            urls = extract_linkedin_urls(page_text)
            if urls:
                url = urls[0]
                return {
                    "url": url,
                    "profile_name": linkedin_profile_name(url),
                    "source": page_url,
                    "status": "confirmed: published on official website",
                }
        return {
            "url": "",
            "profile_name": "",
            "source": response.url,
            "status": "not found on official website",
        }
    except requests.RequestException as exc:
        return {
            "url": "",
            "profile_name": "",
            "source": website,
            "status": f"official website error: {type(exc).__name__}",
        }


def discover_many(
    organisations: dict[str, str],
    *,
    workers: int = 20,
    timeout: int = 10,
    checkpoint: Path | None = None,
) -> dict[str, dict]:
    results: dict[str, dict] = {}
    if checkpoint and checkpoint.is_file():
        import json

        results = json.loads(checkpoint.read_text(encoding="utf-8"))
    pending = {abn: site for abn, site in organisations.items() if abn not in results}
    completed = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(discover_linkedin, site, timeout): abn for abn, site in pending.items()}
        for future in as_completed(futures):
            abn = futures[future]
            try:
                results[abn] = future.result()
            except Exception as exc:  # noqa: BLE001
                results[abn] = {
                    "url": "",
                    "profile_name": "",
                    "source": pending[abn],
                    "status": f"discovery error: {type(exc).__name__}",
                }
            completed += 1
            if checkpoint and completed % 25 == 0:
                import json

                checkpoint.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    if checkpoint:
        import json

        checkpoint.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    return results


def nonprofit_organisations(workbook_path: Path) -> dict[str, str]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["National non-profits"]
    headers = [cell.value for cell in next(ws.iter_rows(max_row=1))]
    index = {header: i for i, header in enumerate(headers)}
    result: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[index["category"]] or "") != "nonprofit":
            continue
        abn = re.sub(r"\D", "", str(row[index["abn"]] or ""))
        if abn:
            result[abn] = str(row[index["website"]] or "").strip()
    return result


def enrich_workbook(source: Path, destination: Path, results: dict[str, dict]) -> dict:
    wb = load_workbook(source)
    updated = 0
    links = 0
    for ws in wb.worksheets:
        if ws.title == "How to use":
            continue
        headers = [cell.value for cell in ws[1]]
        if "category" not in headers or "abn" not in headers:
            continue
        category_col = headers.index("category") + 1
        abn_col = headers.index("abn") + 1
        for name in LINKEDIN_COLUMNS:
            if name not in headers:
                headers.append(name)
                cell = ws.cell(1, len(headers), name)
                cell.fill = PatternFill("solid", fgColor="D9FF4F")
                cell.font = Font(bold=True, color="14160F")
                cell.alignment = Alignment(vertical="center")
        column_map = {name: headers.index(name) + 1 for name in LINKEDIN_COLUMNS}
        for row_no in range(2, ws.max_row + 1):
            category = str(ws.cell(row_no, category_col).value or "")
            if category != "nonprofit":
                continue
            abn = re.sub(r"\D", "", str(ws.cell(row_no, abn_col).value or ""))
            value = results.get(abn) or {
                "url": "",
                "profile_name": "",
                "source": "",
                "status": "not processed",
            }
            original_url = value.get("url", "")
            safe_url = normalise_linkedin_url(original_url)
            safe_status = value.get("status", "")
            if original_url and not safe_url:
                safe_status = "excluded: generic website-builder or invalid organisation link"
            mapped = {
                "linkedin_company_url": safe_url or "",
                "linkedin_profile_name": linkedin_profile_name(safe_url),
                "linkedin_source": value.get("source", ""),
                "linkedin_status": safe_status,
            }
            for name, content in mapped.items():
                cell = ws.cell(row_no, column_map[name], content)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if name == "linkedin_company_url" and content:
                    cell.hyperlink = content
                    cell.style = "Hyperlink"
            updated += 1
            if safe_url:
                links += 1
        for name, col_no in column_map.items():
            ws.column_dimensions[get_column_letter(col_no)].width = 42 if name != "linkedin_status" else 38

    how = wb["How to use"]
    start = how.max_row + 2
    how.cell(start, 1, "LinkedIn enrichment").font = Font(bold=True, size=14, color="14160F")
    how.cell(start + 1, 1, "Organisation pages are discovered only from official organisation websites. LinkedIn itself was not crawled.")
    how.cell(start + 2, 1, "Blank links mean no organisation LinkedIn URL was found on the official website; uncertain matches were not guessed.")
    how.cell(start + 3, 1, f"Unique non-profits processed: {len(results)}")
    confirmed = sum(bool(normalise_linkedin_url(v.get("url", ""))) for v in results.values())
    how.cell(start + 4, 1, f"Confirmed LinkedIn organisation pages: {confirmed}")
    how.column_dimensions["A"].width = max(how.column_dimensions["A"].width or 0, 130)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    return {
        "unique_nonprofits": len(results),
        "confirmed_links": confirmed,
        "rows_updated": updated,
        "linked_rows": links,
    }
