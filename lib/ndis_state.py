"""Add Australian state and company-people views to the NDIS positions workbook."""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATE_NAMES = {
    "NSW": "New South Wales",
    "VIC": "Victoria",
    "QLD": "Queensland",
    "WA": "Western Australia",
    "SA": "South Australia",
    "TAS": "Tasmania",
    "NT": "Northern Territory",
    "ACT": "Australian Capital Territory",
}
STATE_ABBREV_RE = re.compile(r"\b(NSW|VIC|QLD|WA|SA|TAS|NT|ACT)\b", re.I)
POSTCODE_RE = re.compile(r"\b(\d{4})\b")
SOURCE_SHEETS = (
    "NT disability companies",
    "NT non-profits",
    "Darwin disability cos",
    "Darwin non-profits",
    "National disability cos",
    "National non-profits",
)


def state_from_postcode(postcode: str) -> str:
    try:
        value = int(postcode)
    except (TypeError, ValueError):
        return ""
    if 800 <= value <= 999:
        return "NT"
    if 200 <= value <= 299 or 2600 <= value <= 2618 or 2900 <= value <= 2920:
        return "ACT"
    if 1000 <= value <= 1999 or 2000 <= value <= 2599 or 2619 <= value <= 2899 or 2921 <= value <= 2999:
        return "NSW"
    if 3000 <= value <= 3999 or 8000 <= value <= 8999:
        return "VIC"
    if 4000 <= value <= 4999 or 9000 <= value <= 9999:
        return "QLD"
    if 5000 <= value <= 5999:
        return "SA"
    if 6000 <= value <= 6797 or 6800 <= value <= 6999:
        return "WA"
    if 7000 <= value <= 7999:
        return "TAS"
    return ""


def parse_location(address: str) -> dict:
    text = re.sub(r"\s+", " ", str(address or "")).strip()
    abbrev = STATE_ABBREV_RE.search(text)
    postcode_match = POSTCODE_RE.search(text)
    postcode = postcode_match.group(1) if postcode_match else ""
    state = abbrev.group(1).upper() if abbrev else state_from_postcode(postcode)
    suburb = ""
    if text:
        suburb = text.split(",")[0].strip()
        if STATE_ABBREV_RE.fullmatch(suburb) or suburb in {"AU", "Australia"}:
            suburb = ""
    return {
        "state": state,
        "state_name": STATE_NAMES.get(state, ""),
        "suburb": suburb,
        "postcode": postcode,
        "head_office": text,
    }


def _norm_abn(value) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _headers(sheet) -> list:
    return [cell.value for cell in sheet[1]]


def _style(sheet, widths: dict[str, int] | None = None) -> None:
    fill = PatternFill("solid", fgColor="D9FF4F")
    sheet.freeze_panes = "A2"
    if sheet.max_row and sheet.max_column:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="14160F")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = widths or {}
    for number, header in enumerate(_headers(sheet), 1):
        sheet.column_dimensions[get_column_letter(number)].width = widths.get(header, 22)


def _read_sheet_maps(workbook) -> tuple[dict[str, dict], list[dict]]:
    companies: dict[str, dict] = {}
    for title in SOURCE_SHEETS:
        if title not in workbook.sheetnames:
            continue
        sheet = workbook[title]
        headers = _headers(sheet)
        index = {header: number for number, header in enumerate(headers, 1)}
        if "abn" not in index:
            continue
        for row in sheet.iter_rows(min_row=2, values_only=False):
            values = {header: row[number - 1].value for header, number in index.items()}
            abn = _norm_abn(values.get("abn"))
            if not abn:
                continue
            location = parse_location(values.get("head_office") or "")
            current = companies.get(abn)
            if current is None:
                companies[abn] = {
                    "abn": abn,
                    "business_name": str(values.get("business_name") or "").strip(),
                    "legal_name": str(values.get("legal_name") or "").strip(),
                    "website": str(values.get("website") or "").strip(),
                    "category": str(values.get("category") or "").strip(),
                    "has_nt_presence": str(values.get("has_nt_presence") or "").strip(),
                    "has_darwin_presence": str(values.get("has_darwin_presence") or "").strip(),
                    **location,
                }
            else:
                if location["state"] and not current.get("state"):
                    current.update(location)
                if values.get("has_nt_presence") and not current.get("has_nt_presence"):
                    current["has_nt_presence"] = "yes"
                if values.get("has_darwin_presence") and not current.get("has_darwin_presence"):
                    current["has_darwin_presence"] = "yes"
    people: list[dict] = []
    if "People" in workbook.sheetnames:
        sheet = workbook["People"]
        headers = _headers(sheet)
        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            abn = _norm_abn(row.get("abn"))
            if not abn or not str(row.get("person_name") or "").strip():
                continue
            company = companies.get(abn) or {}
            people.append(
                {
                    "abn": abn,
                    "business_name": row.get("business_name") or company.get("business_name") or "",
                    "legal_name": row.get("legal_name") or company.get("legal_name") or "",
                    "person_name": str(row.get("person_name") or "").strip(),
                    "role": str(row.get("role") or "").strip(),
                    "linkedin_profile_url": row.get("linkedin_profile_url") or "",
                    "source_type": row.get("source_type") or "",
                    "source_url": row.get("source_url") or "",
                    "confidence": row.get("confidence") or "",
                    "state": company.get("state") or "",
                    "state_name": company.get("state_name") or "",
                    "head_office": company.get("head_office") or "",
                    "website": company.get("website") or "",
                }
            )
    return companies, people


def _ensure_column(sheet, header: str, after: str | None = None) -> int:
    headers = _headers(sheet)
    if header in headers:
        return headers.index(header) + 1
    if after and after in headers:
        column = headers.index(after) + 2
    else:
        column = sheet.max_column + 1
    sheet.insert_cols(column)
    sheet.cell(1, column, header)
    return column


def _fill_lookup_column(sheet, header: str, after: str, lookup: dict[str, dict], field: str) -> None:
    column = _ensure_column(sheet, header, after)
    headers = _headers(sheet)
    if "abn" not in headers:
        return
    abn_col = headers.index("abn") + 1
    for row_no in range(2, sheet.max_row + 1):
        abn = _norm_abn(sheet.cell(row_no, abn_col).value)
        sheet.cell(row_no, column, (lookup.get(abn) or {}).get(field) or "")


def _append_how_to_use(sheet, captured_at: str, summary: dict) -> None:
    sheet.append(())
    sheet.append(("Company state and people", ""))
    sheet.append(("Updated at (UTC)", captured_at))
    sheet.append(("Companies with a parsed Australian state", summary["companies_with_state"]))
    sheet.append(("Published people joined to a company state", summary["people"]))
    sheet.append(("State column", "Head-office state from the NDIS Commission register address."))
    sheet.append(("Company + people", "One row per published person, with that company's state."))
    sheet.append(("Companies by state", "Every company, its state, and the people found on its official website."))
    sheet.append(
        (
            "Coverage",
            "People are only those published on official websites or verified public results. A blank people list does not mean the company has no staff.",
        )
    )


def write_state_people_workbook(source: Path, destination: Path) -> dict:
    workbook = load_workbook(source)
    companies, people = _read_sheet_maps(workbook)
    people_by_abn: defaultdict[str, list[dict]] = defaultdict(list)
    for person in people:
        people_by_abn[person["abn"]].append(person)
    captured = datetime.now(timezone.utc).replace(microsecond=0).isoformat()

    for title in SOURCE_SHEETS:
        if title not in workbook.sheetnames:
            continue
        _fill_lookup_column(workbook[title], "state", "head_office", companies, "state")
        _fill_lookup_column(workbook[title], "state_name", "state", companies, "state_name")

    if "Organizations" in workbook.sheetnames:
        sheet = workbook["Organizations"]
        _fill_lookup_column(sheet, "state", "legal_name", companies, "state")
        _fill_lookup_column(sheet, "state_name", "state", companies, "state_name")
        _fill_lookup_column(sheet, "head_office", "state_name", companies, "head_office")

    for title in ("People", "Disability Role People", "Disability Job Openings", "Company Role Summary"):
        if title not in workbook.sheetnames:
            continue
        after = "legal_name" if "legal_name" in _headers(workbook[title]) else "business_name"
        _fill_lookup_column(workbook[title], "state", after, companies, "state")
        _fill_lookup_column(workbook[title], "head_office", "state", companies, "head_office")

    for title in ("State summary", "Companies by state", "Company + people"):
        if title in workbook.sheetnames:
            del workbook[title]

    state_sheet = workbook.create_sheet("State summary", 1)
    state_sheet.append(
        (
            "state",
            "state_name",
            "companies",
            "companies_with_published_people",
            "published_people",
            "notes",
        )
    )
    people_counts = Counter(person["state"] or "Unknown" for person in people)
    company_counts = Counter((item.get("state") or "Unknown") for item in companies.values())
    companies_with_people = Counter(
        (companies.get(abn) or {}).get("state") or "Unknown" for abn in people_by_abn
    )
    for state in list(STATE_NAMES) + ["Unknown"]:
        if not company_counts.get(state) and not people_counts.get(state):
            continue
        state_sheet.append(
            (
                state,
                STATE_NAMES.get(state, "State not published"),
                company_counts.get(state, 0),
                companies_with_people.get(state, 0),
                people_counts.get(state, 0),
                "Head-office state from the official NDIS register",
            )
        )
    _style(
        state_sheet,
        {
            "state_name": 28,
            "notes": 46,
            "companies": 16,
            "companies_with_published_people": 32,
            "published_people": 20,
        },
    )

    company_sheet = workbook.create_sheet("Companies by state", 2)
    company_sheet.append(
        (
            "state",
            "state_name",
            "abn",
            "business_name",
            "legal_name",
            "head_office",
            "website",
            "people_count",
            "people_working",
            "has_nt_presence",
            "has_darwin_presence",
            "category",
        )
    )
    for item in sorted(
        companies.values(),
        key=lambda row: (
            row.get("state") or "ZZ",
            str(row.get("business_name") or "").lower(),
        ),
    ):
        staff = people_by_abn.get(item["abn"], [])
        names = "; ".join(
            f"{person['person_name']}" + (f" ({person['role']})" if person.get("role") else "")
            for person in staff
        )
        company_sheet.append(
            (
                item.get("state") or "",
                item.get("state_name") or "",
                item["abn"],
                item.get("business_name") or "",
                item.get("legal_name") or "",
                item.get("head_office") or "",
                item.get("website") or "",
                len(staff),
                names,
                item.get("has_nt_presence") or "",
                item.get("has_darwin_presence") or "",
                item.get("category") or "",
            )
        )
    _style(
        company_sheet,
        {
            "business_name": 36,
            "legal_name": 36,
            "head_office": 36,
            "website": 36,
            "people_working": 60,
            "state_name": 26,
        },
    )

    join_sheet = workbook.create_sheet("Company + people", 3)
    join_sheet.append(
        (
            "state",
            "state_name",
            "abn",
            "business_name",
            "legal_name",
            "head_office",
            "website",
            "person_name",
            "role",
            "linkedin_profile_url",
            "source_type",
            "source_url",
            "confidence",
        )
    )
    for person in sorted(
        people,
        key=lambda row: (
            row.get("state") or "ZZ",
            str(row.get("business_name") or "").lower(),
            str(row.get("person_name") or "").lower(),
        ),
    ):
        join_sheet.append(
            (
                person.get("state") or "",
                person.get("state_name") or "",
                person.get("abn") or "",
                person.get("business_name") or "",
                person.get("legal_name") or "",
                person.get("head_office") or "",
                person.get("website") or "",
                person.get("person_name") or "",
                person.get("role") or "",
                person.get("linkedin_profile_url") or "",
                person.get("source_type") or "",
                person.get("source_url") or "",
                person.get("confidence") or "",
            )
        )
    _style(
        join_sheet,
        {
            "business_name": 36,
            "legal_name": 36,
            "head_office": 32,
            "person_name": 28,
            "role": 36,
            "linkedin_profile_url": 40,
            "source_url": 40,
            "website": 36,
        },
    )

    if "How to use" in workbook.sheetnames:
        _append_how_to_use(
            workbook["How to use"],
            captured,
            {
                "companies_with_state": sum(1 for item in companies.values() if item.get("state")),
                "people": len(people),
            },
        )

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return {
        "companies": len(companies),
        "companies_with_state": sum(1 for item in companies.values() if item.get("state")),
        "people": len(people),
        "companies_with_people": len(people_by_abn),
        "states": dict(company_counts),
        "xlsx": str(destination),
        "captured_at": captured,
    }
