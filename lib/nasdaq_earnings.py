"""Upcoming Nasdaq earnings calendar with third-party technical indicators.

Earnings metadata comes from Nasdaq. Technical values come from TradingView's
robots-allowed global scanner endpoint. Scenario ranges are deterministic,
indicator-based estimates and are not price targets or investment advice.
"""

from __future__ import annotations

import json
import math
import re
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

NASDAQ_EARNINGS_API = "https://api.nasdaq.com/api/calendar/earnings"
NASDAQ_EARNINGS_PAGE = "https://www.nasdaq.com/market-activity/earnings"
TRADINGVIEW_SCANNER = "https://scanner.tradingview.com/global/scan"
TRADINGVIEW_ROBOTS = "https://scanner.tradingview.com/robots.txt"
USER_AGENT = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/139 Safari/537.36"

TV_COLUMNS = (
    "name",
    "description",
    "type",
    "subtype",
    "exchange",
    "country",
    "close",
    "change",
    "volume",
    "market_cap_basic",
    "price_52_week_high",
    "price_52_week_low",
    "RSI",
    "RSI[1]",
    "MACD.macd",
    "MACD.signal",
    "SMA20",
    "SMA50",
    "SMA200",
    "EMA20",
    "EMA50",
    "EMA200",
    "Recommend.All",
    "Recommend.MA",
    "Recommend.Other",
    "ATR",
    "Perf.1M",
    "Perf.3M",
    "Perf.6M",
    "BB.upper",
    "BB.lower",
    "Stoch.K",
    "Stoch.D",
    "ADX",
    "ADX+DI",
    "ADX-DI",
    "CCI20",
    "Mom",
    "VWMA",
    "relative_volume_10d_calc",
)

COLUMNS = (
    "earnings_date",
    "report_time",
    "symbol",
    "company_name",
    "exchange",
    "country",
    "fiscal_quarter_ending",
    "consensus_eps_forecast",
    "number_of_estimates",
    "last_year_report_date",
    "last_year_eps",
    "eps_yoy_change_pct",
    "earnings_expectation",
    "nasdaq_market_cap",
    "price",
    "daily_change_pct",
    "volume",
    "technical_market_cap",
    "52_week_low",
    "52_week_high",
    "distance_from_52w_low_pct",
    "distance_from_52w_high_pct",
    "rsi_14",
    "rsi_signal",
    "macd",
    "macd_signal_line",
    "macd_direction",
    "sma_20",
    "sma_50",
    "sma_200",
    "ema_20",
    "ema_50",
    "ema_200",
    "trend_structure",
    "bollinger_lower",
    "bollinger_upper",
    "bollinger_position",
    "stochastic_k",
    "stochastic_d",
    "adx",
    "adx_plus_di",
    "adx_minus_di",
    "cci_20",
    "momentum",
    "vwma",
    "atr_14",
    "atr_pct",
    "relative_volume_10d",
    "return_1m_pct",
    "return_3m_pct",
    "return_6m_pct",
    "technical_rating_score",
    "technical_rating",
    "bullish_bearish_signal",
    "moving_average_rating",
    "oscillator_rating",
    "future_bias_not_forecast",
    "one_month_scenario_low",
    "one_month_scenario_high",
    "scenario_range_pct",
    "data_status",
    "earnings_source",
    "technical_source",
    "captured_at_utc",
)

HEADER_FILL = PatternFill("solid", fgColor="D9FF4F")
BUY_FILL = PatternFill("solid", fgColor="C6EFCE")
SELL_FILL = PatternFill("solid", fgColor="FFC7CE")
NEUTRAL_FILL = PatternFill("solid", fgColor="FFEB9C")
HEADER_FONT = Font(bold=True, color="14160F")
TITLE_FONT = Font(bold=True, size=15, color="14160F")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="C5C9B8"),
    right=Side(style="thin", color="C5C9B8"),
    top=Side(style="thin", color="C5C9B8"),
    bottom=Side(style="thin", color="C5C9B8"),
)


def parse_number(value: object) -> float | None:
    text = str(value or "").strip()
    if not text or text.lower() in {"n/a", "na", "--", "none"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.+-]", "", text)
    if not cleaned or cleaned in {"+", "-", "."}:
        return None
    try:
        number = float(cleaned)
    except ValueError:
        return None
    return -abs(number) if negative else number


def report_time_label(value: str | None) -> str:
    low = (value or "").lower()
    if "pre-market" in low:
        return "Pre-market"
    if "after-hours" in low:
        return "After hours"
    return "Time not supplied"


def rating_label(score: float | None) -> str:
    if score is None:
        return "Unavailable"
    if score >= 0.5:
        return "Strong Buy"
    if score >= 0.1:
        return "Buy"
    if score > -0.1:
        return "Neutral"
    if score > -0.5:
        return "Sell"
    return "Strong Sell"


def rsi_label(rsi: float | None) -> str:
    if rsi is None:
        return "Unavailable"
    if rsi >= 70:
        return "Overbought"
    if rsi <= 30:
        return "Oversold"
    return "Neutral"


def trend_structure(price: float | None, sma20: float | None, sma50: float | None, sma200: float | None) -> str:
    if None in {price, sma20, sma50, sma200}:
        return "Unavailable"
    if price > sma20 > sma50 > sma200:
        return "Bullish alignment"
    if price < sma20 < sma50 < sma200:
        return "Bearish alignment"
    if price > sma50 and price > sma200:
        return "Positive / mixed"
    if price < sma50 and price < sma200:
        return "Negative / mixed"
    return "Mixed"


def bollinger_position(price: float | None, lower: float | None, upper: float | None) -> str:
    if None in {price, lower, upper} or upper == lower:
        return "Unavailable"
    if price >= upper:
        return "At/above upper band"
    if price <= lower:
        return "At/below lower band"
    position = (price - lower) / (upper - lower)
    if position >= 0.67:
        return "Upper third"
    if position <= 0.33:
        return "Lower third"
    return "Middle third"


def technical_bias(score: float | None, trend: str, macd_direction: str, rsi: float | None) -> str:
    if score is None:
        return "Insufficient technical data"
    points = 0
    points += 2 if score >= 0.5 else 1 if score >= 0.1 else -2 if score <= -0.5 else -1 if score <= -0.1 else 0
    points += 1 if "Bullish" in trend or "Positive" in trend else -1 if "Bearish" in trend or "Negative" in trend else 0
    points += 1 if macd_direction == "Bullish" else -1 if macd_direction == "Bearish" else 0
    if rsi is not None:
        points += -1 if rsi >= 75 else 1 if rsi <= 25 else 0
    if points >= 3:
        return "Bullish technical bias"
    if points <= -3:
        return "Bearish technical bias"
    return "Neutral / mixed technical bias"


def bullish_bearish_signal(bias: str) -> str:
    if bias == "Bullish technical bias":
        return "Bullish"
    if bias == "Bearish technical bias":
        return "Bearish"
    return "Neutral"


def scenario_range(price: float | None, atr: float | None, trading_days: int = 21) -> tuple[float | None, float | None, float | None]:
    if price is None or atr is None or price <= 0 or atr < 0:
        return None, None, None
    move = atr * math.sqrt(trading_days)
    low = max(0.0, price - move)
    high = price + move
    return low, high, move / price * 100


def eps_expectation(forecast: float | None, last_year: float | None) -> tuple[float | None, str]:
    if forecast is None or last_year is None:
        return None, "Insufficient EPS comparison"
    if last_year == 0:
        return None, "Forecast profit" if forecast > 0 else "Forecast loss"
    change = (forecast - last_year) / abs(last_year) * 100
    if last_year < 0 <= forecast:
        label = "Forecast turnaround to profit"
    elif last_year >= 0 > forecast:
        label = "Forecast swing to loss"
    elif change >= 10:
        label = "Forecast EPS growth"
    elif change <= -10:
        label = "Forecast EPS decline"
    else:
        label = "Forecast EPS broadly flat"
    return change, label


def fetch_earnings_day(day: date, session: requests.Session | None = None) -> list[dict]:
    session = session or requests.Session()
    response = session.get(
        NASDAQ_EARNINGS_API,
        params={"date": day.isoformat()},
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/json, text/plain, */*",
            "Origin": "https://www.nasdaq.com",
            "Referer": NASDAQ_EARNINGS_PAGE,
        },
        timeout=45,
    )
    response.raise_for_status()
    payload = response.json()
    return ((payload.get("data") or {}).get("rows") or [])


def fetch_earnings_range(start: date, end: date, cache: Path | None = None, delay: float = 0.25) -> list[dict]:
    if cache and cache.is_file():
        return json.loads(cache.read_text(encoding="utf-8"))
    session = requests.Session()
    output: list[dict] = []
    current = start
    while current <= end:
        for row in fetch_earnings_day(current, session=session):
            output.append({**row, "earningsDate": current.isoformat()})
        current += timedelta(days=1)
        if current <= end and delay:
            time.sleep(delay)
    if cache:
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_text(json.dumps(output, indent=2, ensure_ascii=False), encoding="utf-8")
    return output


def fetch_technicals(symbols: list[str]) -> dict[str, dict]:
    # Nasdaq calendars include NYSE, Nasdaq and American listings.
    exchanges = ("NASDAQ", "NYSE", "AMEX", "OTC")
    tickers = [f"{exchange}:{symbol.replace('/', '.')}" for symbol in symbols for exchange in exchanges]
    payload = {"symbols": {"tickers": tickers, "query": {"types": []}}, "columns": list(TV_COLUMNS)}
    response = requests.post(
        TRADINGVIEW_SCANNER,
        json=payload,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/json",
            "Origin": "https://www.tradingview.com",
            "Referer": "https://www.tradingview.com/",
        },
        timeout=90,
    )
    response.raise_for_status()
    output: dict[str, dict] = {}
    for item in response.json().get("data") or []:
        values = dict(zip(TV_COLUMNS, item.get("d") or []))
        symbol = str(values.get("name") or item.get("s", "").split(":")[-1]).upper()
        candidate = {**values, "_ticker": item.get("s")}
        # Keep the first common-stock match; replace a non-stock match if needed.
        current = output.get(symbol)
        if current is None or (current.get("type") != "stock" and candidate.get("type") == "stock"):
            output[symbol] = candidate
    return output


def build_rows(earnings: list[dict], technicals: dict[str, dict], captured_at: str | None = None) -> list[dict]:
    captured_at = captured_at or datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    rows: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for item in earnings:
        symbol = str(item.get("symbol") or "").strip().upper()
        day = str(item.get("earningsDate") or "")
        if not symbol or (day, symbol) in seen:
            continue
        seen.add((day, symbol))
        tech = technicals.get(symbol) or {}
        price = parse_number(tech.get("close"))
        high52 = parse_number(tech.get("price_52_week_high"))
        low52 = parse_number(tech.get("price_52_week_low"))
        rsi = parse_number(tech.get("RSI"))
        macd = parse_number(tech.get("MACD.macd"))
        macd_signal = parse_number(tech.get("MACD.signal"))
        sma20 = parse_number(tech.get("SMA20"))
        sma50 = parse_number(tech.get("SMA50"))
        sma200 = parse_number(tech.get("SMA200"))
        bb_upper = parse_number(tech.get("BB.upper"))
        bb_lower = parse_number(tech.get("BB.lower"))
        atr = parse_number(tech.get("ATR"))
        rating_score = parse_number(tech.get("Recommend.All"))
        trend = trend_structure(price, sma20, sma50, sma200)
        macd_direction = (
            "Unavailable"
            if macd is None or macd_signal is None
            else "Bullish"
            if macd > macd_signal
            else "Bearish"
            if macd < macd_signal
            else "Neutral"
        )
        scenario_low, scenario_high, scenario_pct = scenario_range(price, atr)
        eps_forecast = parse_number(item.get("epsForecast"))
        last_eps = parse_number(item.get("lastYearEPS"))
        eps_change, expectation = eps_expectation(eps_forecast, last_eps)
        source_page = f"{NASDAQ_EARNINGS_PAGE}?date={day}"
        tv_ticker = tech.get("_ticker") or ""
        technical_source = (
            f"https://www.tradingview.com/symbols/{str(tv_ticker).replace(':', '-')}/" if tv_ticker else ""
        )
        rows.append(
            {
                "earnings_date": day,
                "report_time": report_time_label(item.get("time")),
                "symbol": symbol,
                "company_name": item.get("name") or tech.get("description") or "",
                "exchange": tech.get("exchange") or "",
                "country": tech.get("country") or "",
                "fiscal_quarter_ending": item.get("fiscalQuarterEnding") or "",
                "consensus_eps_forecast": eps_forecast,
                "number_of_estimates": parse_number(item.get("noOfEsts")),
                "last_year_report_date": item.get("lastYearRptDt") or "",
                "last_year_eps": last_eps,
                "eps_yoy_change_pct": eps_change,
                "earnings_expectation": expectation,
                "nasdaq_market_cap": parse_number(item.get("marketCap")),
                "price": price,
                "daily_change_pct": parse_number(tech.get("change")),
                "volume": parse_number(tech.get("volume")),
                "technical_market_cap": parse_number(tech.get("market_cap_basic")),
                "52_week_low": low52,
                "52_week_high": high52,
                "distance_from_52w_low_pct": ((price / low52 - 1) * 100 if price and low52 else None),
                "distance_from_52w_high_pct": ((price / high52 - 1) * 100 if price and high52 else None),
                "rsi_14": rsi,
                "rsi_signal": rsi_label(rsi),
                "macd": macd,
                "macd_signal_line": macd_signal,
                "macd_direction": macd_direction,
                "sma_20": sma20,
                "sma_50": sma50,
                "sma_200": sma200,
                "ema_20": parse_number(tech.get("EMA20")),
                "ema_50": parse_number(tech.get("EMA50")),
                "ema_200": parse_number(tech.get("EMA200")),
                "trend_structure": trend,
                "bollinger_lower": bb_lower,
                "bollinger_upper": bb_upper,
                "bollinger_position": bollinger_position(price, bb_lower, bb_upper),
                "stochastic_k": parse_number(tech.get("Stoch.K")),
                "stochastic_d": parse_number(tech.get("Stoch.D")),
                "adx": parse_number(tech.get("ADX")),
                "adx_plus_di": parse_number(tech.get("ADX+DI")),
                "adx_minus_di": parse_number(tech.get("ADX-DI")),
                "cci_20": parse_number(tech.get("CCI20")),
                "momentum": parse_number(tech.get("Mom")),
                "vwma": parse_number(tech.get("VWMA")),
                "atr_14": atr,
                "atr_pct": (atr / price * 100 if atr is not None and price else None),
                "relative_volume_10d": parse_number(tech.get("relative_volume_10d_calc")),
                "return_1m_pct": parse_number(tech.get("Perf.1M")),
                "return_3m_pct": parse_number(tech.get("Perf.3M")),
                "return_6m_pct": parse_number(tech.get("Perf.6M")),
                "technical_rating_score": rating_score,
                "technical_rating": rating_label(rating_score),
                "bullish_bearish_signal": bullish_bearish_signal(
                    technical_bias(rating_score, trend, macd_direction, rsi)
                ),
                "moving_average_rating": rating_label(parse_number(tech.get("Recommend.MA"))),
                "oscillator_rating": rating_label(parse_number(tech.get("Recommend.Other"))),
                "future_bias_not_forecast": technical_bias(rating_score, trend, macd_direction, rsi),
                "one_month_scenario_low": scenario_low,
                "one_month_scenario_high": scenario_high,
                "scenario_range_pct": scenario_pct,
                "data_status": "Complete" if tech else "Technical data unavailable",
                "earnings_source": source_page,
                "technical_source": technical_source,
                "captured_at_utc": captured_at,
            }
        )
    rows.sort(key=lambda row: (row["earnings_date"], row["report_time"], -(row["nasdaq_market_cap"] or 0)))
    return rows


def us_market_cap_rows(rows: list[dict]) -> list[dict]:
    """Keep US-domiciled companies and order largest to smallest."""
    filtered = [row for row in rows if row.get("country") == "United States"]
    return sorted(
        filtered,
        key=lambda row: (
            -(row.get("technical_market_cap") or row.get("nasdaq_market_cap") or 0),
            row.get("symbol") or "",
        ),
    )


def _style_table(ws: Worksheet, rows: list[dict]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(rows) + 1)}"
    for col_no, header in enumerate(COLUMNS, 1):
        cell = ws.cell(1, col_no, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        cell.border = THIN
    for row_no, row in enumerate(rows, 2):
        for col_no, key in enumerate(COLUMNS, 1):
            cell = ws.cell(row_no, col_no, row.get(key))
            cell.alignment = WRAP
            cell.border = THIN
            if key in {"earnings_source", "technical_source"} and row.get(key):
                cell.hyperlink = row[key]
                cell.style = "Hyperlink"
            if key == "bullish_bearish_signal":
                cell.fill = (
                    BUY_FILL
                    if row.get(key) == "Bullish"
                    else SELL_FILL
                    if row.get(key) == "Bearish"
                    else NEUTRAL_FILL
                )
            if key in {"price", "52_week_low", "52_week_high", "sma_20", "sma_50", "sma_200", "ema_20", "ema_50", "ema_200", "bollinger_lower", "bollinger_upper", "atr_14", "one_month_scenario_low", "one_month_scenario_high"}:
                cell.number_format = "$#,##0.00"
            elif key in {"nasdaq_market_cap", "technical_market_cap", "volume"}:
                cell.number_format = "#,##0"
            elif key.endswith("_pct") or key.startswith("return_") or key.startswith("distance_"):
                cell.number_format = "0.00"
            elif key in {"rsi_14", "stochastic_k", "stochastic_d", "adx", "adx_plus_di", "adx_minus_di", "cci_20", "relative_volume_10d"}:
                cell.number_format = "0.00"
            elif key in {"macd", "macd_signal_line", "momentum", "technical_rating_score"}:
                cell.number_format = "0.0000"
            elif key == "number_of_estimates":
                cell.number_format = "0"
        ws.row_dimensions[row_no].height = 34
    widths = {
        "earnings_date": 13,
        "report_time": 16,
        "symbol": 11,
        "company_name": 34,
        "country": 16,
        "earnings_expectation": 26,
        "trend_structure": 21,
        "technical_rating": 16,
        "bullish_bearish_signal": 18,
        "future_bias_not_forecast": 26,
        "rsi_14": 12,
        "macd": 16,
        "macd_signal_line": 18,
        "stochastic_k": 15,
        "stochastic_d": 15,
        "adx_plus_di": 14,
        "adx_minus_di": 14,
        "technical_rating_score": 20,
        "data_status": 24,
        "earnings_source": 35,
        "technical_source": 35,
        "captured_at_utc": 22,
    }
    for col_no, key in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_no)].width = widths.get(key, 15)
    if rows:
        score_col = COLUMNS.index("technical_rating_score") + 1
        ws.conditional_formatting.add(
            f"{get_column_letter(score_col)}2:{get_column_letter(score_col)}{len(rows) + 1}",
            ColorScaleRule(start_type="num", start_value=-1, start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFEB84", end_type="num", end_value=1, end_color="63BE7B"),
        )


def _how_to_use(
    ws: Worksheet,
    start: date,
    end: date,
    rows: list[dict],
    scope_description: str = "All companies on the Nasdaq earnings calendar",
    order_description: str = "Earnings date, report time, then market cap",
) -> None:
    complete = sum(row["data_status"] == "Complete" for row in rows)
    lines = [
        ("Upcoming earnings + technical indicators", TITLE_FONT),
        ("", None),
        (f"Window: {start.isoformat()} through {end.isoformat()} (30 calendar days). Companies: {len(rows)}.", None),
        (f"Scope: {scope_description}. Order: {order_description}.", None),
        (f"Technical data matched: {complete}; unmatched: {len(rows) - complete}. Missing values are left blank.", None),
        ("", None),
        ("Sources", HEADER_FONT),
        ("Earnings dates, report timing, EPS consensus and prior-year EPS: Nasdaq Earnings Calendar.", None),
        ("Prices and technical indicators: TradingView global scanner (/global/scan is explicitly allowed by its robots.txt).", None),
        ("Nasdaq calendar extraction is limited to this requested 30-day window; the workbook records the public calendar page URL for each date.", None),
        ("", None),
        ("Important limitations", HEADER_FONT),
        ("This is research, not financial advice. Earnings dates and estimates can change. Prices may be delayed.", None),
        ("Technical ratings are snapshots. They do not predict earnings surprises, gaps or future returns.", None),
        ("Future bias is a rules-based summary of TradingView rating, moving-average structure, MACD and RSI.", None),
        ("One-month scenario range = price ± ATR(14) × √21. It is a volatility scenario, not a target or probability interval.", None),
        ("EPS YoY change compares consensus forecast with last year's reported EPS; negative/near-zero denominators can distort percentages.", None),
        ("", None),
        ("Indicator guide", HEADER_FONT),
        ("RSI: ≥70 overbought; ≤30 oversold. MACD: bullish when MACD exceeds its signal line.", None),
        ("SMA/EMA: 20, 50 and 200 sessions. Bullish alignment means price > SMA20 > SMA50 > SMA200.", None),
        ("ADX measures trend strength; Bollinger position shows where price sits within the bands; relative volume compares with 10-day volume.", None),
        ("TradingView rating scores range approximately from -1 (sell) to +1 (buy).", None),
    ]
    ws.column_dimensions["A"].width = 130
    for row_no, (text, font) in enumerate(lines, 1):
        cell = ws.cell(row_no, 1, text)
        cell.alignment = WRAP
        if font:
            cell.font = font
        ws.row_dimensions[row_no].height = 22 if text else 10


def write_workbook(
    path: Path,
    rows: list[dict],
    start: date,
    end: date,
    *,
    scope_description: str = "All companies on the Nasdaq earnings calendar",
    order_description: str = "Earnings date, report time, then market cap",
) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "How to use"
    _how_to_use(guide, start, end, rows, scope_description, order_description)
    all_ws = wb.create_sheet("All upcoming earnings")
    _style_table(all_ws, rows)
    for label, predicate in (
        ("Forecast profit EPS", lambda row: (row["consensus_eps_forecast"] or 0) > 0),
        ("Forecast loss EPS", lambda row: row["consensus_eps_forecast"] is not None and row["consensus_eps_forecast"] < 0),
        ("Bullish signal", lambda row: row["bullish_bearish_signal"] == "Bullish"),
        ("Bearish signal", lambda row: row["bullish_bearish_signal"] == "Bearish"),
        ("Neutral signal", lambda row: row["bullish_bearish_signal"] == "Neutral"),
        ("Pre-market", lambda row: row["report_time"] == "Pre-market"),
        ("After hours", lambda row: row["report_time"] == "After hours"),
        ("Missing technicals", lambda row: row["data_status"] != "Complete"),
    ):
        ws = wb.create_sheet(label)
        _style_table(ws, [row for row in rows if predicate(row)])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path: Path, rows: list[dict]) -> None:
    import csv

    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(COLUMNS))
        writer.writeheader()
        writer.writerows(rows)
