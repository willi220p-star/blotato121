"""ARRCS Darwin teams and positions (roles only, no named people)."""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

FEED_NAME = "arrcs-darwin-teams"
SCOPED_AT = "2026-08-28"
ARRCS_ORIGIN = "https://arrcs.org.au"
JOBS_PORTAL = "https://apply.arrcs.org.au/arrcs/"
SEEK_SEARCH = "https://www.seek.com.au/jobs?keywords=ARRCS%20Darwin"

COLUMNS = (
    "timeframe",
    "team",
    "position",
    "location",
    "status",
    "source_url",
    "notes",
)

DARWIN_HINTS = (
    "darwin",
    "palmerston",
    "tiwi",
    "coconut grove",
    "farrar",
    "casuarina",
    "maluka",
    "nightcliff",
    "parap",
    "stuart park",
    "winnellie",
    "harry's place",
    "harrys place",
    "willeroo",
)
NOT_DARWIN_HINTS = (
    "alice springs",
    "katherine",
    "tennant creek",
    "nhulunbuy",
    "mutitjulu",
    "kaltukatjara",
    "docker river",
    "flynn lodge",
    "pukatja",
    "apy lands",
    "old timers",
    "rocky ridge",
    "yutjuwala",
    "heti perkins",
    "hetti perkins",
)

# Names published on ARRCS pages/PDFs. Export must never include them.
FORBIDDEN_NAME_FRAGMENTS = (
    "jenny messell",
    "fiona stoddart",
    "linda manhire",
    "wendy hubbard",
    "rachel inglis",
    "kylie wells",
    "poppy reece",
    "earanka dhakal",
    "les huddleston",
    "uncle burri",
    "george butler",
    "gucki reissenberger",
    "emelia patterson",
    "avaisha",
    "sarah douglas",
    "irene snell",
    "kay wilson",
    "cheryl",
    "cathy",
    "tito",
    "dave and trevlyn",
)

JOB_TYPE_RE = re.compile(
    r"^(Full Time|Part Time|Casual/Vacation|Apprentice/Trainee|Contract)$",
    re.I,
)
JOB_LOC_RE = re.compile(r"^(.+?)\s*\|\s*(Northern Territory|South Australia|Western Australia|Queensland)$", re.I)

HEADER_FILL = PatternFill("solid", fgColor="D9FF4F")
HEADER_FONT = Font(bold=True, color="14160F")
TITLE_FONT = Font(bold=True, size=14, color="14160F")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(
    left=Side(style="thin", color="C5C9B8"),
    right=Side(style="thin", color="C5C9B8"),
    top=Side(style="thin", color="C5C9B8"),
    bottom=Side(style="thin", color="C5C9B8"),
)


def is_darwin_location(text: str | None) -> bool:
    blob = (text or "").strip().lower()
    if not blob:
        return False
    if any(hint in blob for hint in NOT_DARWIN_HINTS):
        return False
    return any(hint in blob for hint in DARWIN_HINTS)


def contains_person_name(text: str | None) -> bool:
    blob = (text or "").lower()
    return any(name in blob for name in FORBIDDEN_NAME_FRAGMENTS)


def parse_job_board_text(text: str) -> list[dict]:
    """Parse ApplyNow listing text copied from the jobs portal."""
    lines = [re.sub(r"\s+", " ", line).strip() for line in (text or "").splitlines()]
    lines = [line for line in lines if line]
    jobs: list[dict] = []
    i = 0
    while i < len(lines) - 2:
        job_type = lines[i + 1]
        if not JOB_TYPE_RE.match(job_type):
            i += 1
            continue
        title = lines[i]
        location = ""
        classification = ""
        blurb_parts: list[str] = []
        j = i + 2
        while j < len(lines):
            loc_match = JOB_LOC_RE.match(lines[j])
            if loc_match:
                location = loc_match.group(1).strip()
                if j + 1 < len(lines) and "," in lines[j + 1]:
                    classification = lines[j + 1]
                break
            if lines[j] in {"1", "2", "3", "Next", "Previous", "Filters"}:
                break
            blurb_parts.append(lines[j])
            j += 1
        if location:
            jobs.append(
                {
                    "title": title,
                    "job_type": job_type,
                    "location": location,
                    "classification": classification,
                    "blurb": " ".join(blurb_parts).strip(),
                }
            )
            i = j + 1
            continue
        i += 1
    return jobs


def darwin_jobs_from_board(text: str) -> list[dict]:
    return [job for job in parse_job_board_text(text) if is_darwin_location(job.get("location"))]


def _row(
    timeframe: str,
    team: str,
    position: str,
    location: str,
    status: str,
    source_url: str,
    notes: str,
) -> dict:
    row = {
        "timeframe": timeframe,
        "team": team,
        "position": position,
        "location": location,
        "status": status,
        "source_url": source_url,
        "notes": notes,
    }
    leaked = " ".join(row.values())
    if contains_person_name(leaked):
        raise ValueError(f"Row includes a person name: {position} / {notes}")
    if timeframe == "present" and not is_darwin_location(location) and "Darwin" not in location:
        raise ValueError(f"Non-Darwin present row: {location}")
    return row


def researched_rows() -> list[dict]:
    """Darwin / Palmerston teams and positions compiled from ARRCS public pages and reviews."""
    j = "https://arrcs.org.au/facility/juninga-centre-aged-care-facility/"
    tg = "https://arrcs.org.au/facility/terrace-gardens-aged-care-facility/"
    darwin = "https://arrcs.org.au/darwin/"
    rost = "https://arrcs.org.au/kicking-goals-at-darwin-rost/"
    home = "https://arrcs.org.au/help-at-home/"
    allied = "https://arrcs.org.au/allied-health/"
    vol = "https://arrcs.org.au/work-with-arrcs/volunteer/"
    careers = "https://arrcs.org.au/work-with-arrcs/careers/"
    trainees = "https://arrcs.org.au/work-with-arrcs/first-nations-traineeships/"
    exec_url = "https://arrcs.org.au/our-executive-team/"
    biennial_2025 = "https://arrcs.org.au/wp-content/uploads/2026/04/Arrcs-2025-Binenial-review-Final-Version-Compressed.pdf"
    biennial_2023 = "https://arrcs.org.au/wp-content/uploads/2025/10/ARRCS-Biennial-Report-2022-2023-Nov-version-Final.pdf"
    annual_2021 = "https://arrcs.org.au/wp-content/uploads/2025/10/arrcs-annual-report-2021-2.pdf"
    annual_2020 = "https://arrcs.org.au/wp-content/uploads/2025/10/ARRCS-Annual-Review-2020.pdf"
    jenny = "https://arrcs.org.au/why-jenny-loves-working-at-juninga/"
    eoi = "https://arrcs-external.applynow.net.au/jobs/ARRCSEOI-expression-of-interest"

    rows: list[dict] = []

    # --- Present: Juninga ---
    rows += [
        _row("present", "Juninga Centre Aged Care", "Service manager", "Coconut Grove, Darwin", "current", j, "Current facility page names this role. 26 residential beds plus 10 independent-living cabins supported by Community Care Darwin. 24-hour staff and nurse-call."),
        _row("present", "Juninga Centre Aged Care", "Nursing staff (24-hour)", "Coconut Grove, Darwin", "current", j, "Nurse-call, clinical care, respite and palliative care on site."),
        _row("present", "Juninga Centre Aged Care", "Care / personal care workers", "Coconut Grove, Darwin", "current", j, "Day-to-day support for residents; cultural and social programs."),
        _row("present", "Juninga Centre Aged Care", "Leisure and lifestyle / activities staff", "Coconut Grove, Darwin", "current", j, "Planned activities five days a week (music, movies, campfires, fishing, outings)."),
        _row("present", "Juninga Centre Aged Care", "Kitchen / hospitality staff", "Coconut Grove, Darwin", "current", j, "On-site menus including cultural meals cooked on an open fire."),
        _row("present", "Juninga Centre Aged Care", "Residential respite staff", "Coconut Grove, Darwin", "current", j, "Residential respite listed on the current facility page."),
        _row("present", "Juninga Centre Aged Care", "Palliative care staff", "Coconut Grove, Darwin", "current", j, "Special care: palliative care."),
        _row("present", "Juninga Centre Aged Care", "Visiting general practitioner", "Coconut Grove, Darwin", "current (visiting)", j, "Regular GP visits. Visiting clinician, not necessarily an ARRCS employee."),
        _row("present", "Juninga Centre Aged Care", "Visiting allied health (physio, podiatry, OT, speech)", "Coconut Grove, Darwin", "current (visiting)", j, "Listed on the current facility page."),
    ]

    # --- Present: Terrace Gardens (Palmerston / Greater Darwin) ---
    rows += [
        _row("present", "Terrace Gardens Aged Care", "Nursing staff (24-hour)", "Farrar / Palmerston", "current", tg, "88-bed home, three wings, nurse-call, 24-hour staff access."),
        _row("present", "Terrace Gardens Aged Care", "Care / personal care workers", "Farrar / Palmerston", "current", tg, "Permanent residential care across three wings."),
        _row("present", "Terrace Gardens Aged Care", "Memory Support Unit care staff", "Farrar / Palmerston", "current", tg, "Dementia / Memory Support Unit with garden. Careers page confirms this unit is staffed."),
        _row("present", "Terrace Gardens Aged Care", "Leisure and lifestyle / activities staff", "Farrar / Palmerston", "current", tg, "Arts and crafts, pet therapy, music, barbecues and outings."),
        _row("present", "Terrace Gardens Aged Care", "Residential respite staff", "Farrar / Palmerston", "current", tg, "Residential respite listed on the current facility page."),
        _row("present", "Terrace Gardens Aged Care", "Palliative care staff", "Farrar / Palmerston", "current", tg, "Special care: palliative care."),
        _row("present", "Terrace Gardens Aged Care", "Visiting general practitioner", "Farrar / Palmerston", "current (visiting)", tg, "Regular GP visits. Visiting clinician, not necessarily an ARRCS employee."),
        _row("present", "Terrace Gardens Aged Care", "Visiting allied health (physio, podiatry, OT, speech)", "Farrar / Palmerston", "current (visiting)", tg, "Listed on the current facility page."),
    ]

    # --- Present: Darwin ROST ---
    rows += [
        _row("present", "Darwin ROST (Respite Options for Senior Territorians)", "ROST worker", "Suburban Darwin respite house", "current", rost, "Respite house in a suburban Darwin home. Workers support goal-based respite and daily activities."),
        _row("present", "Darwin ROST (Respite Options for Senior Territorians)", "Lifestyle / activities staff", "Suburban Darwin respite house", "current", rost, "Lifestyle calendar including entertainers, puzzles, colouring, reading and outings."),
    ]

    # --- Present: Community Care Darwin / Help at Home ---
    rows += [
        _row("present", "Community Care Darwin / Help at Home", "Home care / personal care worker", "Darwin (in-home + Juninga ILU cabins)", "current", home, "Darwin page lists in-home care. Juninga cabins are supported by Community Care Darwin. 2022–23 review lists Darwin Home Care as a Darwin service."),
        _row("present", "Community Care Darwin / Help at Home", "Home care nurse", "Darwin (in-home)", "current", home, "Nurses visit consumers at home for clinical care and medication. Help at Home lists nursing and medication management."),
        _row("present", "Community Care Darwin / Help at Home", "In-home respite worker", "Darwin (in-home)", "current", home, "In-home respite: lifestyle, light duties, personal care, shopping and meal preparation."),
        _row("present", "Community Care Darwin / Help at Home", "Home maintenance / modification worker", "Darwin (in-home)", "current", home, "Help at Home: home maintenance and modification, cleaning and housework, linen."),
        _row("present", "Community Care Darwin / Help at Home", "Transport / driver", "Darwin", "current", home, "Help at Home lists transport. Volunteer drivers also cover this work."),
        _row("present", "Community Care Darwin / Help at Home", "Meals / food services", "Darwin", "current", home, "Help at Home lists meals."),
        _row("present", "Community Care Darwin / Help at Home", "NDIS / disability support worker", "Darwin", "current", darwin, "Darwin enquiry form includes disability supports. 2022–23 review includes a Darwin NDIS participant pathway."),
        _row("present", "Community Care Darwin / Help at Home", "Dementia care worker", "Darwin (in-home)", "current", home, "Help at Home lists dementia care and management."),
        _row("present", "Community Care Darwin / Help at Home", "Palliative care (home)", "Darwin (in-home)", "current", home, "Help at Home lists palliative care and carers support."),
    ]

    # --- Present: Maluka (Palmerston) ---
    rows += [
        _row("present", "Maluka Day Programs", "Centre-based day respite / day program staff", "Palmerston", "current", biennial_2025, "2025 biennial: overnight respite opened at Maluka in Palmerston in 2022–23; centre-based and overnight respite remain integral. 2022–23 review lists Maluka as Darwin-region centre-based day respite."),
        _row("present", "Maluka Day Programs", "Overnight respite staff", "Palmerston", "current", biennial_2025, "Overnight respite opened at Maluka in the 2022–23 financial year."),
    ]

    # --- Present: Allied health delivered in Darwin ---
    rows += [
        _row("present", "Allied health (Darwin delivery)", "Occupational therapist", "Darwin homes and aged-care facilities", "current", allied, "Territory-wide allied health; Darwin is a service/enquiry location. Delivered in home, facilities or centres."),
        _row("present", "Allied health (Darwin delivery)", "Physiotherapist", "Darwin homes and aged-care facilities", "current", allied, "Listed on allied-health and Darwin facility pages."),
        _row("present", "Allied health (Darwin delivery)", "Podiatrist", "Darwin homes and aged-care facilities", "current", allied, "Listed on allied-health and Darwin facility pages."),
        _row("present", "Allied health (Darwin delivery)", "Speech therapist", "Darwin homes and aged-care facilities", "current", allied, "Listed on allied-health and Darwin facility pages."),
        _row("present", "Allied health (Darwin delivery)", "Hydrotherapy", "Darwin (Help at Home / allied health)", "current", allied, "Hydrotherapy is listed on Help at Home and allied health. Confirm site before outreach — Fred McKay Day Therapy Centre is Alice Springs."),
    ]

    # --- Present: Volunteers ---
    rows += [
        _row("present", "Volunteers (Darwin)", "Bus / car driving volunteer", "Darwin (location option on the form)", "current", vol, "Volunteer page is always recruiting. Darwin is a location option."),
        _row("present", "Volunteers (Darwin)", "Deliveries volunteer", "Darwin (location option on the form)", "current", vol, "Assisting with deliveries."),
        _row("present", "Volunteers (Darwin)", "Arts, craft, music and storytelling volunteer", "Darwin (location option on the form)", "current", vol, "Activity volunteer roles."),
        _row("present", "Volunteers (Darwin)", "Companion volunteer", "Darwin (location option on the form)", "current", vol, "Friendly companion role."),
    ]

    # --- Present: Darwin HQ / NT teams based at Harry's Place ---
    rows += [
        _row("present", "Executive / corporate (Harry's Place)", "Chief Executive Officer", "Tiwi, Darwin (Harry's Place, 1 Willeroo Street)", "current", exec_url, "Executive page lists CEO. 2025 biennial: CEO replaced the General Manager title during the independence transition. Role only — no names in this workbook."),
        _row("present", "Executive / corporate (Harry's Place)", "Executive Manager of Corporate Services", "Tiwi, Darwin (Harry's Place)", "current", biennial_2025, "Created as part of the 2024–25 finance/payroll independence from UnitingCare Queensland."),
        _row("present", "Executive / corporate (Harry's Place)", "Finance and payroll", "Tiwi, Darwin (Harry's Place)", "current", biennial_2025, "Finance and payroll moved onto ARRCS-managed systems in 2024–early 2025."),
        _row("present", "Clinical Education Team", "Clinical educator / graduate-nurse preceptor", "NT team; works with Darwin facilities", "current", biennial_2025, "Supports the Gerontology Graduate Nursing Program (12–18 month pathway). Face-to-face study days. Not Darwin-only, but Darwin facilities are in scope."),
        _row("present", "First Nations Programs", "First Nations Programs team member", "Darwin (Juninga and Terrace Gardens in scope)", "current", biennial_2025, "2025 biennial: team supports Juninga Women's Focus Groups and Darwin-region cultural work. Do not treat Alice Springs / remote-only programs as Darwin."),
        _row("present", "RAP Working Group", "RAP Working Group member", "Darwin HQ plus site representatives", "current", biennial_2025, "Reactivated in the 2024–25 snapshot. NT-wide group with Darwin HQ involvement."),
        _row("present", "Gerontology Graduate Nursing Program", "Graduate nurse", "Darwin facilities in an NT-wide program", "current", biennial_2025, "Pilot 2023; first cohort progressing 12–18 months with Clinical Education Team support."),
    ]

    # --- Past: superseded titles, closed intakes, historically documented Darwin roles ---
    rows += [
        _row("past", "Executive / corporate (Harry's Place)", "General Manager", "Tiwi, Darwin (Harry's Place)", "superseded", biennial_2025, "Replaced by Chief Executive Officer during the 2024–25 independence transition."),
        _row("past", "Executive / corporate (Harry's Place)", "Project manager (systems independence)", "Darwin HQ", "time-limited", biennial_2025, "Dedicated project-manager role oversaw finance/payroll cutover scope, timeline, workforce alignment and change (2024–early 2025)."),
        _row("past", "Juninga Centre Aged Care", "Service manager / clinical nursing leadership", "Coconut Grove, Darwin", "historical (role still exists)", jenny, "Feature article documents 23 years in this Darwin service-manager / nursing leadership role and a 2019 Territory nursing award. Present sheet already lists Service manager — this row is the historical evidence."),
        _row("past", "Terrace Gardens Aged Care", "Clinical nurse manager", "Farrar / Palmerston", "last documented 2020", annual_2020, "2020 annual review describes a Clinical Nurse Manager pathway at Terrace Gardens. Not named on the current facility page."),
        _row("past", "Terrace Gardens Aged Care", "Activities officer", "Farrar / Palmerston", "last documented 2022–23", biennial_2023, "2022–23 biennial lists an Activities Officer at Terrace Gardens."),
        _row("past", "Terrace Gardens Aged Care", "Team leader, First Nations", "Farrar / Palmerston", "last documented 2022–23", biennial_2023, "2022–23 biennial lists a First Nations Team Leader at Terrace Gardens."),
        _row("past", "Facility leadership (Darwin / Palmerston)", "Facility service manager", "Darwin and Palmerston", "last documented 2022–23", biennial_2023, "2022–23 biennial lists Facility Service Managers for Darwin, Palmerston and Katherine."),
        _row("past", "Clinical leadership", "Regional manager, clinical", "Darwin and Palmerston (also other NT sites)", "last documented 2022–23", biennial_2023, "2022–23 biennial: Regional Manager Clinical covering Darwin and Palmerston, working with Royal Darwin and Palmerston GPs, plus other NT sites."),
        _row("past", "Community Care Darwin / Help at Home", "Regional manager, community (Home Care and NDIS)", "Darwin (jobs expo also ran in other NT centres)", "last documented 2022–23", biennial_2023, "Led Home Care and NDIS recruiting at Darwin and other NT jobs expos."),
        _row("past", "Community Care Darwin / Help at Home", "Senior support worker", "Darwin Community Care", "last documented 2022–23", biennial_2023, "Named as a Darwin Community Care role in the 2022–23 biennial (role title only retained here)."),
        _row("past", "First Nations Programs", "Aboriginal liaison officer", "Darwin (covering Darwin and Katherine)", "last documented 2020", annual_2020, "Role commenced March 2020; a permanent Darwin-based role covering Darwin and Katherine."),
        _row("past", "First Nations Programs", "First Nations cultural specialist", "Juninga and Terrace Gardens, Darwin / Palmerston", "last documented 2021", annual_2021, "2021 annual review: cultural specialist rotated through Juninga and Terrace Gardens plus other NT aged-care homes."),
        _row("past", "First Nations traineeships", "Personal care worker trainee (Cert III Individual Support)", "Darwin was a program location", "closed intake (2024)", trainees, "2024 intake closed. One intake a year; Darwin remains a form location for the next round (see Future)."),
        _row("past", "First Nations traineeships", "Hospitality worker trainee (Cert III Commercial Cookery)", "Darwin was a program location", "closed intake (2024)", trainees, "2024 intake closed."),
        _row("past", "First Nations traineeships", "Leisure and lifestyle officer trainee (Cert IV Leisure and Health)", "Darwin was a program location", "closed intake (2024)", trainees, "2024 intake closed."),
        _row("past", "First Nations traineeships", "IT worker trainee (Cert III IT)", "Darwin was a program location", "closed intake (2024)", trainees, "2024 intake closed."),
        _row("past", "Recruitment", "Expression of interest (generic)", "All locations including Darwin", "closed 21 Feb 2022", eoi, "ApplyNow EOI job ARRCSEOI. Page states this job has closed. Not a live Darwin vacancy."),
        _row("past", "Return to Country / Troopy Program", "Troopy program trainee", "NT program; Darwin mentioned in 2021 review", "last documented 2020–21", annual_2021, "2021 review documents a 2020 Troopy Program trainee pathway. Darwin appears in the same regional story; not Darwin-only."),
        _row("past", "Juninga Centre Aged Care", "Refurbishment / capital works team", "Coconut Grove, Darwin", "completed", biennial_2025, "2024–25 snapshot: Juninga refurbishments completed and opened. Event/project, not an ongoing FTE line."),
        _row("past", "Maluka Day Programs", "Maluka Day Programs opening team", "Palmerston", "opened 2022–23", biennial_2023, "Opening of Maluka Day Programs in Palmerston recorded in the 2022–23 biennial."),
    ]

    # --- Future: live Darwin ads + pipeline ---
    rows += [
        _row(
            "future",
            "First Nations traineeships / administration",
            "First Nations trainee — administration",
            "Darwin",
            "live vacancy",
            "https://apply.arrcs.org.au/arrcs/9792357/",
            "Live on apply.arrcs.org.au (scraped 28 Aug 2026). Apprentice/Trainee. Must identify as Aboriginal or Torres Strait Islander. Classification: Administration & Office Support.",
        ),
        _row(
            "future",
            "Executive / residential services",
            "Executive manager — residential services",
            "Darwin (NT-wide remit, 12 facilities)",
            "live vacancy",
            "https://apply.arrcs.org.au/arrcs/9784188/",
            "Live on apply.arrcs.org.au (scraped 28 Aug 2026). Full-time. Based Darwin; leads residential aged care across the Territory.",
        ),
        _row("future", "First Nations traineeships", "Personal care worker trainee (Cert III Individual Support)", "Darwin (form location)", "next intake / enquire", trainees, "Program still advertised. 2024 intake closed; page says one intake a year. Darwin is a location on the enquiry form."),
        _row("future", "First Nations traineeships", "Hospitality worker trainee (Cert III Commercial Cookery)", "Darwin (form location)", "next intake / enquire", trainees, "Paid on-the-job training with CDU TAFE. Next intake via enquiry form."),
        _row("future", "First Nations traineeships", "Leisure and lifestyle officer trainee (Cert IV Leisure and Health)", "Darwin (form location)", "next intake / enquire", trainees, "Training via Dovaston Training and Assessment Centre. Next intake via enquiry form."),
        _row("future", "First Nations traineeships", "IT worker trainee (Cert III IT)", "Darwin (form location)", "next intake / enquire", trainees, "Training via CDU TAFE. Next intake via enquiry form."),
        _row("future", "Volunteers (Darwin)", "Volunteer (all advertised volunteer types)", "Darwin", "always recruiting", vol, "Volunteer page: always looking. Darwin is a location option."),
        _row("future", "Careers pipeline (not a named Darwin ad)", "Nursing (registered / enrolled)", "Darwin is named as a Territory workplace", "always recruiting", careers, "Careers page lists nursing among ongoing career types from tropical Darwin to the Central Desert. Not a live Darwin vacancy on the jobs board (28 Aug 2026)."),
        _row("future", "Careers pipeline (not a named Darwin ad)", "Home care assistance", "Darwin is named as a Territory workplace", "always recruiting", careers, "Careers page lists home care assistance as an ongoing career type. No live Darwin home-care ad on the jobs board (28 Aug 2026)."),
        _row(
            "future",
            "SEEK search (blocked)",
            "(none listed — search not readable from this environment)",
            "Darwin keywords",
            "blocked",
            SEEK_SEARCH,
            "SEEK ?keywords=ARRCS Darwin is allowed by our SEEK policy, but this environment gets Cloudflare 403. Individual SEEK /job/ URLs are never fetched. Do not invent SEEK vacancies.",
        ),
    ]

    return rows


def merge_live_jobs(rows: list[dict], live_jobs: list[dict] | None) -> list[dict]:
    """Overlay freshly scraped Darwin jobs onto the researched future sheet."""
    if not live_jobs:
        return rows
    out = [row for row in rows if not (row["timeframe"] == "future" and row["status"] == "live vacancy")]
    for job in live_jobs:
        href = job.get("url") or job.get("href") or JOBS_PORTAL
        if href.startswith("/"):
            href = "https://apply.arrcs.org.au" + href
        out.append(
            _row(
                "future",
                job.get("team") or "ARRCS jobs portal",
                job.get("title") or job.get("position") or "Vacancy",
                job.get("location") or "Darwin",
                "live vacancy",
                href,
                job.get("notes")
                or " ".join(
                    part
                    for part in (
                        job.get("job_type"),
                        job.get("classification"),
                        job.get("blurb"),
                    )
                    if part
                ),
            )
        )
    return out


def assert_export_safe(rows: list[dict]) -> None:
    for row in rows:
        blob = " ".join(str(v) for v in row.values())
        if contains_person_name(blob):
            raise ValueError(f"Person name leaked in {row.get('position')}")
        if row["timeframe"] not in {"past", "present", "future"}:
            raise ValueError(f"Bad timeframe {row['timeframe']}")


def _style_header(ws: Worksheet, headers: tuple[str, ...]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        cell.border = THIN


def _write_rows(ws: Worksheet, rows: list[dict]) -> None:
    _style_header(ws, COLUMNS)
    for r_idx, row in enumerate(rows, 2):
        for c_idx, key in enumerate(COLUMNS, 1):
            cell = ws.cell(r_idx, c_idx, row.get(key) or "")
            cell.alignment = WRAP
            cell.border = THIN
    widths = {"A": 12, "B": 42, "C": 46, "D": 38, "E": 28, "F": 52, "G": 62}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 22
    for r_idx in range(2, len(rows) + 2):
        ws.row_dimensions[r_idx].height = 48


def _how_to_use(ws: Worksheet, counts: dict[str, int]) -> None:
    lines = [
        ("ARRCS Darwin — teams and positions (roles only)", TITLE_FONT),
        ("", None),
        (f"Scoped {SCOPED_AT}. Darwin / Palmerston / Tiwi / Coconut Grove / Farrar / Casuarina / Maluka only.", None),
        ("No named people. Job ads go on Future. Family Support (Mutitjulu / Alice Springs) and Flynn Lodge are excluded.", None),
        ("", None),
        ("Sheet", None),
        (f"Present — {counts.get('present', 0)} current Darwin teams/roles from live pages and current reviews", None),
        (f"Past — {counts.get('past', 0)} superseded titles, closed intakes, and historically documented Darwin roles", None),
        (f"Future — {counts.get('future', 0)} live Darwin vacancies, next traineeship intake, volunteer/careers pipeline", None),
        (f"All in this list — {counts.get('all', 0)} rows", None),
        ("Sources and gaps — robots, jobs portal, SEEK miss, what was excluded", None),
        ("", None),
        ("Live Darwin vacancies on apply.arrcs.org.au (23 jobs NT-wide; 2 located Darwin):", None),
        ("1. First Nations Trainee — Administration  https://apply.arrcs.org.au/arrcs/9792357/", None),
        ("2. Executive Manager — Residential Services  https://apply.arrcs.org.au/arrcs/9784188/", None),
        ("", None),
        ("HQ: PO Box 43021 Casuarina NT 0811 · Harry's Place, 1 Willeroo Street, Tiwi NT 0810 · 08 8982 5200", None),
        ("Juninga: 113 Dick Ward Drive, Coconut Grove · Terrace Gardens: 1 Kettle Street, Farrar NT 0830", None),
    ]
    ws.column_dimensions["A"].width = 118
    for idx, (text, font) in enumerate(lines, 1):
        cell = ws.cell(idx, 1, text)
        cell.alignment = WRAP
        if font:
            cell.font = font
        ws.row_dimensions[idx].height = 20 if text else 10


def _sources_sheet(ws: Worksheet) -> None:
    rows = [
        ("Item", "Detail"),
        ("robots.txt", "https://arrcs.org.au/robots.txt — User-agent: * Disallow: (empty). Crawl allowed. Sitemap listed."),
        ("Darwin hub", "https://arrcs.org.au/darwin/"),
        ("Juninga", "https://arrcs.org.au/facility/juninga-centre-aged-care-facility/"),
        ("Terrace Gardens", "https://arrcs.org.au/facility/terrace-gardens-aged-care-facility/"),
        ("Darwin ROST", "https://arrcs.org.au/kicking-goals-at-darwin-rost/"),
        ("Help at Home", "https://arrcs.org.au/help-at-home/"),
        ("Allied health", "https://arrcs.org.au/allied-health/"),
        ("Volunteer", "https://arrcs.org.au/work-with-arrcs/volunteer/ (plural /volunteers/ is 404)"),
        ("Careers", "https://arrcs.org.au/work-with-arrcs/careers/"),
        ("Traineeships", "https://arrcs.org.au/work-with-arrcs/first-nations-traineeships/"),
        ("Executive", "https://arrcs.org.au/our-executive-team/ — titles only used"),
        ("Jobs portal", "https://apply.arrcs.org.au/arrcs/ — StealthyFetcher, 23 jobs, paginated. 2 Darwin."),
        ("SEEK", f"{SEEK_SEARCH} — Cloudflare 403 from this environment. /job/ never fetched."),
        ("LinkedIn", "Not crawled."),
        ("Excluded", "Alice Springs, Katherine, Tennant Creek, Nhulunbuy, Mutitjulu, Kaltukatjara, Flynn Lodge, Family Support childcare/nutrition, PALM (not listed as Darwin)."),
        ("L.I.F.E. model", "2025 biennial pilot is Flynn Lodge (Alice Springs) — not listed as a Darwin future team."),
        ("People", "No person names in any data sheet."),
    ]
    _style_header(ws, ("Item", "Detail"))
    for r_idx, (item, detail) in enumerate(rows[1:], 2):
        ws.cell(r_idx, 1, item).alignment = WRAP
        ws.cell(r_idx, 2, detail).alignment = WRAP
        for c in (1, 2):
            ws.cell(r_idx, c).border = THIN
        ws.row_dimensions[r_idx].height = 36
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100


def write_workbook(path: Path, rows: list[dict] | None = None) -> dict:
    rows = list(rows or researched_rows())
    assert_export_safe(rows)
    present = [r for r in rows if r["timeframe"] == "present"]
    past = [r for r in rows if r["timeframe"] == "past"]
    future = [r for r in rows if r["timeframe"] == "future"]
    counts = {"present": len(present), "past": len(past), "future": len(future), "all": len(rows)}

    wb = Workbook()
    how = wb.active
    how.title = "How to use"
    _how_to_use(how, counts)
    _write_rows(wb.create_sheet("Present"), present)
    _write_rows(wb.create_sheet("Past"), past)
    _write_rows(wb.create_sheet("Future"), future)
    _write_rows(wb.create_sheet("All in this list"), rows)
    _sources_sheet(wb.create_sheet("Sources and gaps"))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return counts


def rows_to_csv(rows: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(COLUMNS))
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key) or "" for key in COLUMNS})
    return buf.getvalue()


def feed_payload(rows: list[dict], counts: dict[str, int]) -> dict:
    return {
        "feed": FEED_NAME,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "scoped_at": SCOPED_AT,
        "policy": {
            "people": "Roles and teams only. No named people.",
            "geography": "Darwin / Palmerston / Tiwi / Coconut Grove / Farrar / Casuarina / Maluka only.",
            "jobs": "Live job ads are Future. SEEK ?keywords= only; /job/ never fetched.",
            "linkedin": "Not crawled.",
        },
        "counts": counts,
        "records": rows,
    }


def feed_markdown(payload: dict) -> str:
    counts = payload.get("counts") or {}
    lines = [
        f"# {payload.get('feed')}",
        "",
        f"Scoped {payload.get('scoped_at')}. {counts.get('all', 0)} Darwin roles "
        f"(present {counts.get('present', 0)} · past {counts.get('past', 0)} · future {counts.get('future', 0)}).",
        "",
        "Roles and teams only. No named people. Darwin / Palmerston only.",
        "",
        "## Counts",
        f"- present: {counts.get('present', 0)}",
        f"- past: {counts.get('past', 0)}",
        f"- future: {counts.get('future', 0)}",
        "",
        "## Live Darwin vacancies",
    ]
    for row in payload.get("records") or []:
        if row.get("status") == "live vacancy":
            lines.append(f"- {row['position']} — {row['location']} — {row['source_url']}")
    return "\n".join(lines) + "\n"


def write_outputs(feeds_dir: Path, rows: list[dict] | None = None) -> dict:
    rows = list(rows or researched_rows())
    xlsx = feeds_dir / "ARRCS_Darwin_teams_roles.xlsx"
    counts = write_workbook(xlsx, rows)
    payload = feed_payload(rows, counts)
    (feeds_dir / "arrcs-darwin-teams.json").write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    (feeds_dir / "arrcs-darwin-teams.md").write_text(feed_markdown(payload), encoding="utf-8")
    (feeds_dir / "arrcs-darwin-teams.csv").write_text(rows_to_csv(rows), encoding="utf-8")
    payload["xlsx"] = str(xlsx)
    return payload
