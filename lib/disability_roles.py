"""Filter disability-care roles and discover published job openings."""

from __future__ import annotations

import json
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlparse, urlunparse

import requests
from lxml import html
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lib.website_people import (
    USER_AGENT,
    _clean_text,
    _robots_allows,
    _session,
    organisations_from_workbook,
)

CORPORATE_EXCLUSIONS = re.compile(
    r"\b(?:account|administration|administrative|business analysis|"
    r"communications?|compliance|customer service|events?|finance|financial|"
    r"fleet|fundraising|human resources|information technology|marketing|"
    r"office|payroll|project|quality|recruitment|sales|technology)\b",
    re.I,
)
ROLE_RULES: tuple[tuple[str, re.Pattern], ...] = (
    (
        "Disability Support Worker",
        re.compile(
            r"\b(?:disability|community|residential|mental health|complex care|"
            r"aged care|personal care|home care|youth|outreach)?\s*support worker\b|"
            r"\bsupport professional\b",
            re.I,
        ),
    ),
    (
        "Support Coordinator",
        re.compile(
            r"\b(?:specialist\s+)?support coord(?:inator|ination)\b|"
            r"\bcoordinator of supports\b",
            re.I,
        ),
    ),
    ("Case Manager / Case Worker", re.compile(r"\bcase\s*(?:manager|worker)\b", re.I)),
    ("Occupational Therapist", re.compile(r"\boccupational therap(?:ist|y)\b|\bOT\b", re.I)),
    (
        "Behaviour Support Practitioner",
        re.compile(
            r"\bbehaviou?r support practitioner\b|\bpositive behaviou?r support\b|"
            r"\bPBS practitioner\b",
            re.I,
        ),
    ),
    ("Speech Pathologist", re.compile(r"\bspeech (?:pathologist|therapist|pathology)\b", re.I)),
    ("Physiotherapist", re.compile(r"\bphysiotherap(?:ist|y)\b", re.I)),
    (
        "Psychologist",
        re.compile(r"\b(?:clinical |registered |provisional )?psychologist\b", re.I),
    ),
    (
        "Social Worker",
        re.compile(r"\b(?:mental health |accredited )?social worker\b", re.I),
    ),
    (
        "Psychosocial Recovery Coach",
        re.compile(r"\b(?:psychosocial )?recovery coach\b", re.I),
    ),
    (
        "Personal Care Worker / Carer",
        re.compile(
            r"\bpersonal care (?:worker|assistant)\b|\bcare(?:r|giver| worker)\b|"
            r"\bhome care worker\b",
            re.I,
        ),
    ),
    (
        "Allied Health Professional",
        re.compile(
            r"\ballied health\b|\bdevelopmental educator\b|\bexercise physiologist\b",
            re.I,
        ),
    ),
    (
        "Nurse",
        re.compile(
            r"\b(?:registered|enrolled|clinical|community|disability) nurse\b|\bRN\b",
            re.I,
        ),
    ),
    (
        "Therapy Assistant",
        re.compile(r"\b(?:allied health|occupational therapy|therapy) assistant\b", re.I),
    ),
    (
        "Employment Consultant",
        re.compile(
            r"\b(?:disability )?employment (?:consultant|advisor|coach|specialist)\b",
            re.I,
        ),
    ),
    (
        "Counsellor",
        re.compile(r"\b(?:mental health |rehabilitation )?counsell?or\b", re.I),
    ),
    (
        "Disability / Community Practitioner",
        re.compile(
            r"\b(?:disability|community|mental health|clinical) practitioner\b",
            re.I,
        ),
    ),
    (
        "Service / House Manager",
        re.compile(
            r"\b(?:disability|care|service|services|residential|house|accommodation|"
            r"supported independent living|SIL) manager\b",
            re.I,
        ),
    ),
    (
        "Team Leader",
        re.compile(r"\b(?:community |clinical |care |support |house )?team leader\b|\bhouse leader\b", re.I),
    ),
    (
        "Disability / Service Coordinator",
        re.compile(
            r"\b(?:care|service delivery|services|clinical|intake|program|client|"
            r"recovery|community|disability|NDIS|roster(?:ing)?) coordinator\b",
            re.I,
        ),
    ),
    (
        "Other Disability Care Role",
        re.compile(
            r"\b(?:community|peer|youth|mental health|residential|lifestyle) worker\b|"
            r"\bdisability support\b|\btherapy aide\b|\blifestyle assistant\b",
            re.I,
        ),
    ),
)
CAREER_HINTS = (
    "career",
    "careers",
    "jobs",
    "job",
    "vacancies",
    "vacancy",
    "work-with-us",
    "join-us",
    "join-our-team",
    "employment",
    "opportunities",
)
ATS_HOSTS = (
    "ethicaljobs.com.au",
    "seek.com.au",
    "livehire.com",
    "employmenthero.com",
    "workdayjobs.com",
    "bamboohr.com",
    "smartjobs.qld.gov.au",
)
JOB_CONTEXT_RE = re.compile(
    r"\b(?:apply|career|careers|employment|job|jobs|join|opening|openings|"
    r"opportunit(?:y|ies)|position|positions|recruit|recruitment|role|roles|"
    r"vacancy|vacancies|work with us)\b",
    re.I,
)


def classify_disability_role(title: str) -> str:
    value = _clean_text(title)
    if not value:
        return ""
    for category, pattern in ROLE_RULES:
        if not pattern.search(value):
            continue
        if category in {"Team Leader", "Disability / Service Coordinator"} and CORPORATE_EXCLUSIONS.search(value):
            return ""
        return category
    return ""


def filtered_people(path: Path) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["People"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    rows: list[dict] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        category = classify_disability_role(str(row.get("role") or ""))
        if category:
            rows.append({"role_category": category, **row})
    return rows


def _has_job_context(target_url: str, text: str) -> bool:
    parsed = urlparse(target_url)
    value = f"{parsed.path.replace('-', ' ').replace('_', ' ')} {text}"
    return bool(JOB_CONTEXT_RE.search(value)) or any(
        parsed.netloc.lower().endswith(host) for host in ATS_HOSTS
    )


def _career_pages(home_url: str, content: str, limit: int = 6) -> tuple[list[str], list[dict]]:
    parsed = urlparse(home_url)
    try:
        document = html.fromstring(content)
    except (TypeError, ValueError):
        return [home_url], []
    candidates: list[tuple[int, str]] = []
    linked_openings: list[dict] = []
    for anchor in document.xpath("//a[@href]"):
        href = anchor.get("href") or ""
        target = urljoin(home_url, href)
        target_parsed = urlparse(target)
        text = _clean_text(" ".join(anchor.itertext()))
        low = f"{target_parsed.path} {text}".lower()
        category = classify_disability_role(text)
        if category and _has_job_context(target, text):
            linked_openings.append(
                {
                    "role_category": category,
                    "position_title": text,
                    "job_url": target,
                    "source_page": home_url,
                    "status": "Published link; closing date not supplied",
                    "date_posted": "",
                    "valid_through": "",
                    "location": "",
                    "employment_type": "",
                }
            )
        is_same_host = target_parsed.netloc.lower() == parsed.netloc.lower()
        is_ats = any(target_parsed.netloc.lower().endswith(host) for host in ATS_HOSTS)
        score = sum(hint in low for hint in CAREER_HINTS)
        if not score:
            continue
        if is_same_host:
            clean = urlunparse(
                (target_parsed.scheme, target_parsed.netloc, target_parsed.path, "", "", "")
            )
            candidates.append((score, clean))
        elif is_ats:
            linked_openings.append(
                {
                    "role_category": category,
                    "position_title": text or "External careers listing",
                    "job_url": target,
                    "source_page": home_url,
                    "status": "Official website links to external careers listing",
                    "date_posted": "",
                    "valid_through": "",
                    "location": "",
                    "employment_type": "",
                }
            )
    ordered = [home_url]
    for _, url in sorted(candidates, key=lambda item: (-item[0], len(item[1]))):
        if url not in ordered:
            ordered.append(url)
        if len(ordered) >= limit:
            break
    return ordered, linked_openings


def _walk_json(value):
    if isinstance(value, list):
        for item in value:
            yield from _walk_json(item)
    elif isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk_json(child)


def _location_text(value) -> str:
    if isinstance(value, list):
        return "; ".join(filter(None, (_location_text(item) for item in value)))
    if not isinstance(value, dict):
        return _clean_text(str(value or ""))
    address = value.get("address") or value
    if not isinstance(address, dict):
        return _clean_text(str(address or ""))
    return ", ".join(
        str(address.get(key) or "").strip()
        for key in ("addressLocality", "addressRegion", "addressCountry")
        if address.get(key)
    )


def extract_job_openings(content: str, source_url: str) -> list[dict]:
    try:
        document = html.fromstring(content or "")
    except (TypeError, ValueError):
        return []
    found: list[dict] = []
    for script in document.xpath("//script[@type='application/ld+json']/text()"):
        try:
            payload = json.loads(script)
        except (TypeError, ValueError):
            continue
        for item in _walk_json(payload):
            kinds = item.get("@type")
            kinds = kinds if isinstance(kinds, list) else [kinds]
            if "JobPosting" not in kinds:
                continue
            title = _clean_text(str(item.get("title") or ""))
            category = classify_disability_role(title)
            if not category:
                continue
            valid = str(item.get("validThrough") or "")
            if valid and valid[:10] < date.today().isoformat():
                continue
            found.append(
                {
                    "role_category": category,
                    "position_title": title,
                    "job_url": urljoin(source_url, str(item.get("url") or source_url)),
                    "source_page": source_url,
                    "status": "Published JobPosting",
                    "date_posted": str(item.get("datePosted") or ""),
                    "valid_through": valid,
                    "location": _location_text(item.get("jobLocation")),
                    "employment_type": _clean_text(str(item.get("employmentType") or "")),
                }
            )
    for anchor in document.xpath("//a[@href]"):
        title = _clean_text(" ".join(anchor.itertext()))
        category = classify_disability_role(title)
        target = urljoin(source_url, anchor.get("href") or "")
        if (
            not category
            or not 2 <= len(title.split()) <= 16
            or not _has_job_context(target, title)
        ):
            continue
        found.append(
            {
                "role_category": category,
                "position_title": title,
                "job_url": target,
                "source_page": source_url,
                "status": "Published careers-page link; closing date not supplied",
                "date_posted": "",
                "valid_through": "",
                "location": "",
                "employment_type": "",
            }
        )
    unique: dict[tuple[str, str], dict] = {}
    for opening in found:
        key = (
            opening["job_url"].lower().rstrip("/"),
            re.sub(r"\W", "", opening["position_title"].lower()),
        )
        current = unique.get(key)
        if current is None or current["status"] != "Published JobPosting":
            unique[key] = opening
    return list(unique.values())


def discover_organisation_jobs(website: str, timeout: int = 7) -> dict:
    website = (website or "").strip()
    if not website:
        return {"status": "no official website", "openings": [], "pages_checked": []}
    if not website.startswith(("http://", "https://")):
        website = "https://" + website
    parsed = urlparse(website)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return {"status": "invalid official website", "openings": [], "pages_checked": []}
    if not _robots_allows(website, timeout):
        return {"status": "blocked by robots.txt", "openings": [], "pages_checked": []}
    try:
        response = _session().get(website, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
    except requests.RequestException as exc:
        return {
            "status": f"official website error: {type(exc).__name__}",
            "openings": [],
            "pages_checked": [],
        }
    pages, openings = _career_pages(response.url, response.text)
    checked: list[str] = []
    for index, page_url in enumerate(pages):
        if not _robots_allows(page_url, timeout):
            continue
        content = response.text if index == 0 else ""
        final_url = response.url if index == 0 else page_url
        if index:
            try:
                page_response = _session().get(page_url, timeout=timeout, allow_redirects=True)
                page_response.raise_for_status()
                content = page_response.text
                final_url = page_response.url
            except requests.RequestException:
                continue
        checked.append(final_url)
        openings.extend(extract_job_openings(content, final_url))
    unique: dict[tuple[str, str], dict] = {}
    for opening in openings:
        if not opening.get("role_category"):
            continue
        key = (
            str(opening.get("job_url") or "").lower().rstrip("/"),
            re.sub(r"\W", "", str(opening.get("position_title") or "").lower()),
        )
        unique.setdefault(key, opening)
    return {
        "status": "positions found" if unique else "no filtered positions found",
        "openings": list(unique.values()),
        "pages_checked": checked,
    }


def discover_many_jobs(
    organisations: dict[str, dict],
    *,
    workers: int = 32,
    timeout: int = 7,
    checkpoint: Path | None = None,
) -> dict[str, dict]:
    results: dict[str, dict] = {}
    if checkpoint and checkpoint.is_file():
        results = json.loads(checkpoint.read_text(encoding="utf-8"))
    pending = {abn: item for abn, item in organisations.items() if abn not in results}
    completed = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {
            pool.submit(discover_organisation_jobs, item.get("website", ""), timeout): abn
            for abn, item in pending.items()
        }
        for future in as_completed(futures):
            abn = futures[future]
            try:
                results[abn] = future.result()
            except Exception as exc:  # noqa: BLE001
                results[abn] = {
                    "status": f"discovery error: {type(exc).__name__}",
                    "openings": [],
                    "pages_checked": [],
                }
            completed += 1
            if checkpoint and completed % 50 == 0:
                checkpoint.parent.mkdir(parents=True, exist_ok=True)
                checkpoint.write_text(
                    json.dumps(results, indent=2, ensure_ascii=False),
                    encoding="utf-8",
                )
    if checkpoint:
        checkpoint.parent.mkdir(parents=True, exist_ok=True)
        checkpoint.write_text(json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
    return results


def _style_sheet(sheet, widths: dict[str, int]) -> None:
    fill = PatternFill("solid", fgColor="D9FF4F")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color="14160F")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    headers = [cell.value for cell in sheet[1]]
    for number, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(number)].width = widths.get(header, 22)


def write_role_workbook(
    source: Path,
    destination: Path,
    organisations: dict[str, dict],
    people: list[dict],
    job_results: dict[str, dict],
) -> dict:
    workbook = load_workbook(source)
    for title in ("Disability Role People", "Disability Job Openings", "Company Role Summary"):
        if title in workbook.sheetnames:
            del workbook[title]
    people_sheet = workbook.create_sheet("Disability Role People")
    people_headers = (
        "role_category",
        "abn",
        "business_name",
        "legal_name",
        "person_name",
        "position_title",
        "linkedin_profile_url",
        "source_type",
        "source_url",
        "confidence",
    )
    people_sheet.append(people_headers)
    for row in sorted(
        people,
        key=lambda row: (
            str(row.get("business_name") or "").lower(),
            row["role_category"],
            str(row.get("role") or "").lower(),
            str(row.get("person_name") or "").lower(),
        ),
    ):
        people_sheet.append(
            (
                row["role_category"],
                row.get("abn"),
                row.get("business_name"),
                row.get("legal_name"),
                row.get("person_name"),
                row.get("role"),
                row.get("linkedin_profile_url"),
                row.get("source_type"),
                row.get("source_url"),
                row.get("confidence"),
            )
        )
    opening_sheet = workbook.create_sheet("Disability Job Openings")
    opening_headers = (
        "role_category",
        "abn",
        "business_name",
        "legal_name",
        "position_title",
        "status",
        "date_posted",
        "valid_through",
        "location",
        "employment_type",
        "job_url",
        "source_page",
    )
    opening_sheet.append(opening_headers)
    openings: list[dict] = []
    for abn, result in job_results.items():
        organisation = organisations.get(abn) or {}
        for opening in result.get("openings") or []:
            openings.append({"abn": abn, **organisation, **opening})
    for row in sorted(
        openings,
        key=lambda row: (
            str(row.get("business_name") or "").lower(),
            row["role_category"],
            str(row.get("position_title") or "").lower(),
        ),
    ):
        opening_sheet.append(tuple(row.get(header, "") for header in opening_headers))
    summary_sheet = workbook.create_sheet("Company Role Summary")
    summary_headers = (
        "abn",
        "business_name",
        "legal_name",
        "filtered_role_categories",
        "filtered_position_titles",
        "distinct_position_titles",
        "people_positions_found",
        "published_openings_found",
        "total_position_records",
        "website_research_status",
        "captured_at_utc",
    )
    summary_sheet.append(summary_headers)
    people_by_abn: dict[str, list[dict]] = {}
    openings_by_abn: dict[str, list[dict]] = {}
    for row in people:
        people_by_abn.setdefault(str(row["abn"]), []).append(row)
    for row in openings:
        openings_by_abn.setdefault(str(row["abn"]), []).append(row)
    captured = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    relevant_abns = set(people_by_abn) | set(openings_by_abn)
    for abn in sorted(
        relevant_abns,
        key=lambda value: str((organisations.get(value) or {}).get("business_name") or "").lower(),
    ):
        organisation = organisations.get(abn) or {}
        company_people = people_by_abn.get(abn, [])
        company_openings = openings_by_abn.get(abn, [])
        categories = sorted(
            {row["role_category"] for row in company_people + company_openings}
        )
        titles = sorted(
            {
                _clean_text(str(row.get("role") or row.get("position_title") or ""))
                for row in company_people + company_openings
                if row.get("role") or row.get("position_title")
            },
            key=str.lower,
        )
        summary_sheet.append(
            (
                abn,
                organisation.get("business_name", ""),
                organisation.get("legal_name", ""),
                "; ".join(categories),
                "; ".join(titles),
                len(titles),
                len(company_people),
                len(company_openings),
                len(company_people) + len(company_openings),
                (job_results.get(abn) or {}).get("status", "not processed"),
                captured,
            )
        )
    for sheet in (people_sheet, opening_sheet, summary_sheet):
        _style_sheet(
            sheet,
            {
                "business_name": 34,
                "legal_name": 34,
                "person_name": 26,
                "position_title": 42,
                "filtered_role_categories": 44,
                "filtered_position_titles": 90,
                "linkedin_profile_url": 44,
                "source_url": 48,
                "job_url": 48,
                "source_page": 48,
                "website_research_status": 30,
            },
        )
    for row in people_sheet.iter_rows(min_row=2):
        for column in (7, 9):
            cell = row[column - 1]
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    for row in opening_sheet.iter_rows(min_row=2):
        for column in (11, 12):
            cell = row[column - 1]
            if cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    how = workbook["How to use"]
    start = how.max_row + 2
    notes = (
        "Disability position filtering",
        "Filtered role people are based on public staff titles already linked to each organization by ABN.",
        "Job openings are discovered from official organization websites and career pages; no Apify or LinkedIn crawl is used.",
        "Published links without a closing date may be stale and should be checked before applying.",
        f"Filtered people records: {len(people)}",
        f"Published disability job links: {len(openings)}",
        f"Companies with a filtered people position or opening: {len(relevant_abns)}",
    )
    for offset, note in enumerate(notes):
        cell = how.cell(start + offset, 1, note)
        if offset == 0:
            cell.font = Font(bold=True, size=14, color="14160F")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return {
        "unique_organisations": len(organisations),
        "filtered_people": len(people),
        "published_openings": len(openings),
        "companies_with_filtered_positions": len(relevant_abns),
    }


def build_role_workbook(
    source: Path,
    destination: Path,
    *,
    workers: int = 32,
    timeout: int = 7,
    checkpoint: Path | None = None,
) -> dict:
    organisations = organisations_from_workbook(source)
    people = filtered_people(source)
    jobs = discover_many_jobs(
        organisations,
        workers=workers,
        timeout=timeout,
        checkpoint=checkpoint,
    )
    return write_role_workbook(source, destination, organisations, people, jobs)
