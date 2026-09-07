"""Extract staff published on organisations' official websites.

The crawler never fetches LinkedIn. Individual LinkedIn URLs are accepted only
when an official organisation website publishes the link.
"""

from __future__ import annotations

import html as html_module
import json
import re
import threading
import urllib.robotparser
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urljoin, urlparse, urlunparse

import requests
from lxml import html
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

USER_AGENT = "Mozilla/5.0 (compatible; NDIS-public-staff-discovery/1.0)"
PAGE_HINTS = (
    "team",
    "people",
    "leadership",
    "management",
    "staff",
    "board",
    "directors",
    "governance",
    "about-us",
    "who-we-are",
)
ROLE_WORDS = re.compile(
    r"\b(?:chief|ceo|coo|cfo|director|manager|chair|president|founder|"
    r"secretary|treasurer|executive|officer|coordinator|leader|head|"
    r"practitioner|therapist|psychologist|advisor|consultant|supervisor)\b",
    re.I,
)
NAME_ROLE_LINE = re.compile(
    r"^\s*(?P<name>[A-ZÀ-ÖØ-Þ][A-Za-zÀ-ÖØ-öø-ÿ'.-]+"
    r"(?:\s+(?:[A-ZÀ-ÖØ-Þ][A-Za-zÀ-ÖØ-öø-ÿ'.-]+|de|del|di|da|la|le|van|von|der)){1,4})"
    r"\s*[–—]\s*(?P<role>.{2,100})\s*$"
)
NAME_EXCLUSIONS = {
    "about us",
    "our team",
    "meet the team",
    "meet our team",
    "our people",
    "leadership team",
    "board of directors",
    "management team",
    "contact us",
    "who we are",
    "executive team",
    "share this on",
}
NAME_PARTICLES = {"de", "del", "di", "da", "la", "le", "van", "von", "der"}
NEGATIVE_PAGE_PATHS = (
    "/blog/",
    "/news/",
    "/article/",
    "/articles/",
    "/content-hub/",
    "/resources/",
    "/events/",
)
_local = threading.local()


def _session() -> requests.Session:
    session = getattr(_local, "session", None)
    if session is None:
        session = requests.Session()
        session.headers.update(
            {
                "User-Agent": USER_AGENT,
                "Accept": "text/html,application/xhtml+xml",
            }
        )
        _local.session = session
    return session


def normalise_person_linkedin_url(value: str) -> str:
    value = html_module.unescape((value or "").replace("\\/", "/").strip())
    if value.startswith("//"):
        value = "https:" + value
    parsed = urlparse(value)
    host = parsed.netloc.lower().split(":", 1)[0]
    match = re.match(r"^/in/([^/?#]+)", parsed.path, re.I)
    if (
        host not in {"linkedin.com", "www.linkedin.com"}
        and not host.endswith(".linkedin.com")
    ) or not match:
        return ""
    slug = match.group(1).strip("/")
    if len(slug) < 3 or slug.lower() in {"login", "signup", "undefined"}:
        return ""
    return f"https://www.linkedin.com/in/{slug}"


def _clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", html_module.unescape(value or "")).strip()


def _plausible_name(value: str) -> bool:
    value = _clean_text(value).strip("–—|,;:")
    if not value or value.lower() in NAME_EXCLUSIONS or len(value) > 70:
        return False
    words = value.split()
    if not 2 <= len(words) <= 5:
        return False
    if ROLE_WORDS.search(value):
        return False
    for word in words:
        if word.lower() in NAME_PARTICLES:
            continue
        if not re.match(r"^[A-ZÀ-ÖØ-Þ][A-Za-zÀ-ÖØ-öø-ÿ'.-]*$", word):
            return False
    return True


def _name_from_slug(url: str) -> str:
    slug = urlparse(url).path.rstrip("/").split("/")[-1]
    slug = re.sub(r"-[0-9a-f]{6,}$", "", slug, flags=re.I)
    candidate = re.sub(r"[-_]+", " ", slug).strip().title()
    return candidate if _plausible_name(candidate) else ""


def _nearest_name(node, fallback_url: str = "") -> str:
    candidates = [_clean_text(" ".join(node.itertext()))]
    parent = node
    for _ in range(4):
        parent = parent.getparent()
        if parent is None:
            break
        for heading in parent.xpath(".//h1|.//h2|.//h3|.//h4|.//h5|.//*[@itemprop='name']"):
            candidates.append(_clean_text(" ".join(heading.itertext())))
    for candidate in candidates:
        if _plausible_name(candidate):
            return candidate
    return _name_from_slug(fallback_url)


def _nearest_role(node, person_name: str) -> str:
    following = node.xpath(
        "following::*[self::p or self::span or self::h4 or self::h5][1]"
    )
    if following:
        segment = _clean_text(" ".join(following[0].itertext()))
        if (
            segment
            and segment != person_name
            and len(segment) <= 120
            and ROLE_WORDS.search(segment)
        ):
            return segment
    for sibling in list(node.itersiblings())[:3]:
        if not isinstance(sibling.tag, str):
            continue
        segment = _clean_text(" ".join(sibling.itertext()))
        if (
            segment
            and segment != person_name
            and len(segment) <= 120
            and ROLE_WORDS.search(segment)
        ):
            return segment
    parent = node
    for _ in range(4):
        parent = parent.getparent()
        if parent is None:
            break
        parent_text = _clean_text(" ".join(parent.itertext()))
        if len(parent_text) > 500:
            continue
        segments = [
            _clean_text(" ".join(item.itertext()))
            for item in parent.xpath("./p|./span|./div|./*[@itemprop='jobTitle']")
        ]
        for segment in segments:
            if (
                segment
                and segment != person_name
                and len(segment) <= 120
                and ROLE_WORDS.search(segment)
            ):
                return segment
    return ""


def _linkedin_from_value(value) -> str:
    values = value if isinstance(value, list) else [value]
    for candidate in values:
        if isinstance(candidate, str):
            url = normalise_person_linkedin_url(candidate)
            if url:
                return url
    return ""


def _jsonld_people(document, source_url: str) -> list[dict]:
    found: list[dict] = []

    def walk(value) -> None:
        if isinstance(value, list):
            for item in value:
                walk(item)
            return
        if not isinstance(value, dict):
            return
        kind = value.get("@type")
        kinds = kind if isinstance(kind, list) else [kind]
        if "Person" in kinds:
            name = _clean_text(str(value.get("name") or ""))
            if _plausible_name(name):
                linkedin_url = _linkedin_from_value(value.get("sameAs"))
                found.append(
                    {
                        "person_name": name,
                        "role": _clean_text(str(value.get("jobTitle") or "")),
                        "linkedin_profile_url": linkedin_url,
                        "source_type": "official website structured Person",
                        "source_url": source_url,
                        "confidence": "high" if linkedin_url else "medium",
                    }
                )
        for child in value.values():
            walk(child)

    for script in document.xpath("//script[@type='application/ld+json']/text()"):
        try:
            walk(json.loads(script))
        except (TypeError, ValueError):
            continue
    return found


def extract_people_from_html(content: str, source_url: str) -> list[dict]:
    try:
        document = html.fromstring(content or "")
    except (ValueError, TypeError):
        return []
    for unwanted in document.xpath("//script[not(@type='application/ld+json')]|//style|//nav|//footer"):
        unwanted.drop_tree()
    found = _jsonld_people(document, source_url)
    for anchor in document.xpath("//a[@href]"):
        linkedin_url = normalise_person_linkedin_url(anchor.get("href") or "")
        if not linkedin_url:
            continue
        name = _nearest_name(anchor, linkedin_url)
        if not name:
            continue
        found.append(
            {
                "person_name": name,
                "role": _nearest_role(anchor, name),
                "linkedin_profile_url": linkedin_url,
                "source_type": "official website LinkedIn link",
                "source_url": source_url,
                "confidence": "high",
            }
        )
    for heading in document.xpath("//h2|//h3|//h4|//h5|//*[@itemtype='https://schema.org/Person']//*[@itemprop='name']"):
        name = _clean_text(" ".join(heading.itertext()))
        if not _plausible_name(name):
            continue
        role = _nearest_role(heading, name)
        if not role:
            continue
        found.append(
            {
                "person_name": name,
                "role": role,
                "linkedin_profile_url": "",
                "source_type": "official website staff page",
                "source_url": source_url,
                "confidence": "medium",
            }
        )
    for text_node in document.xpath("//text()[normalize-space()]"):
        line = _clean_text(str(text_node))
        match = NAME_ROLE_LINE.match(line)
        if not match:
            continue
        name = _clean_text(match.group("name"))
        role = _clean_text(match.group("role"))
        if not _plausible_name(name) or not ROLE_WORDS.search(role):
            continue
        found.append(
            {
                "person_name": name,
                "role": role,
                "linkedin_profile_url": "",
                "source_type": "official website staff page",
                "source_url": source_url,
                "confidence": "medium",
            }
        )
    deduped: dict[str, dict] = {}
    for person in found:
        key = re.sub(r"\W", "", person["person_name"].lower())
        current = deduped.get(key)
        if (
            current is None
            or (person["linkedin_profile_url"] and not current["linkedin_profile_url"])
            or (not current["role"] and person["role"])
        ):
            if current and not person["role"]:
                person["role"] = current["role"]
            deduped[key] = person
    return list(deduped.values())


def _robots_allows(url: str, timeout: int) -> bool:
    parsed = urlparse(url)
    robots_url = f"{parsed.scheme}://{parsed.netloc}/robots.txt"
    try:
        response = _session().get(robots_url, timeout=timeout, allow_redirects=True)
        if response.status_code in {401, 403}:
            return False
        if response.status_code >= 400:
            return True
        parser = urllib.robotparser.RobotFileParser()
        parser.set_url(robots_url)
        parser.parse(response.text.splitlines())
        return parser.can_fetch(USER_AGENT, url)
    except requests.RequestException:
        return True


def _candidate_people_pages(home_url: str, content: str, limit: int = 5) -> list[str]:
    parsed = urlparse(home_url)
    candidates: list[tuple[int, str]] = []
    try:
        document = html.fromstring(content)
    except (ValueError, TypeError):
        return [home_url]
    for anchor in document.xpath("//a[@href]"):
        target = urljoin(home_url, anchor.get("href") or "")
        target_parsed = urlparse(target)
        if target_parsed.netloc.lower() != parsed.netloc.lower():
            continue
        if any(part in target_parsed.path.lower() for part in NEGATIVE_PAGE_PATHS):
            continue
        text = f"{target_parsed.path} {_clean_text(' '.join(anchor.itertext()))}".lower()
        path_segments = {
            segment
            for segment in re.split(r"[/_-]+", target_parsed.path.lower())
            if segment
        }
        anchor_text = _clean_text(" ".join(anchor.itertext())).lower()
        score = sum(
            hint in path_segments
            or hint.replace("-", " ") in anchor_text
            or f"/{hint}/" in f"{target_parsed.path.lower().rstrip('/')}/"
            for hint in PAGE_HINTS
        )
        if not score:
            continue
        clean = urlunparse(
            (target_parsed.scheme, target_parsed.netloc, target_parsed.path, "", "", "")
        )
        candidates.append((score, clean))
    ordered = [home_url]
    for _, url in sorted(candidates, key=lambda item: (-item[0], len(item[1]))):
        if url not in ordered:
            ordered.append(url)
        if len(ordered) >= limit:
            break
    return ordered


def discover_website_people(website: str, timeout: int = 8) -> dict:
    website = (website or "").strip()
    if not website:
        return {"status": "no official website", "people": [], "pages_checked": []}
    if not website.startswith(("http://", "https://")):
        website = "https://" + website
    parsed = urlparse(website)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return {"status": "invalid official website", "people": [], "pages_checked": []}
    if not _robots_allows(website, timeout):
        return {"status": "blocked by robots.txt", "people": [], "pages_checked": []}
    try:
        response = _session().get(website, timeout=timeout, allow_redirects=True)
        response.raise_for_status()
    except requests.RequestException as exc:
        return {
            "status": f"official website error: {type(exc).__name__}",
            "people": [],
            "pages_checked": [],
        }
    pages = _candidate_people_pages(response.url, response.text)
    people: list[dict] = []
    checked: list[str] = []
    for index, page_url in enumerate(pages):
        if not _robots_allows(page_url, timeout):
            continue
        page_content = response.text if index == 0 else ""
        final_url = response.url if index == 0 else page_url
        if index:
            try:
                page_response = _session().get(page_url, timeout=timeout, allow_redirects=True)
                page_response.raise_for_status()
                page_content = page_response.text
                final_url = page_response.url
            except requests.RequestException:
                continue
        checked.append(final_url)
        people.extend(extract_people_from_html(page_content, final_url))
    unique: dict[str, dict] = {}
    for person in people:
        key = re.sub(r"\W", "", person["person_name"].lower())
        current = unique.get(key)
        if current is None or (
            person["linkedin_profile_url"] and not current["linkedin_profile_url"]
        ):
            unique[key] = person
    return {
        "status": "people found" if unique else "no public staff found",
        "people": list(unique.values()),
        "pages_checked": checked,
    }


def organisations_from_workbook(path: Path) -> dict[str, dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    organisations: dict[str, dict] = {}
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = list(next(iterator))
        except StopIteration:
            continue
        index = {header: number for number, header in enumerate(headers)}
        if "abn" not in index:
            continue
        for row in iterator:
            abn = re.sub(r"\D", "", str(row[index["abn"]] or ""))
            if not abn:
                continue
            organisation = organisations.setdefault(abn, {"abn": abn})
            for key in (
                "business_name",
                "legal_name",
                "website",
                "linkedin_company_url",
                "linkedin_profile_name",
                "linkedin_source",
                "linkedin_status",
            ):
                if key in index and row[index[key]] and not organisation.get(key):
                    organisation[key] = str(row[index[key]]).strip()
    return organisations


def discover_many_websites(
    organisations: dict[str, dict],
    *,
    workers: int = 24,
    timeout: int = 8,
    checkpoint: Path | None = None,
) -> dict[str, dict]:
    results: dict[str, dict] = {}
    if checkpoint and checkpoint.is_file():
        results = json.loads(checkpoint.read_text(encoding="utf-8"))
    pending = {
        abn: organisation
        for abn, organisation in organisations.items()
        if abn not in results
    }
    completed = 0
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {
            pool.submit(discover_website_people, item.get("website", ""), timeout): abn
            for abn, item in pending.items()
        }
        for future in as_completed(futures):
            abn = futures[future]
            try:
                results[abn] = future.result()
            except Exception as exc:  # noqa: BLE001
                results[abn] = {
                    "status": f"discovery error: {type(exc).__name__}",
                    "people": [],
                    "pages_checked": [],
                }
            completed += 1
            if checkpoint and completed % 50 == 0:
                checkpoint.parent.mkdir(parents=True, exist_ok=True)
                checkpoint.write_text(
                    json.dumps(results, indent=2, ensure_ascii=False),
                    encoding="utf-8",
                )
    if checkpoint:
        checkpoint.parent.mkdir(parents=True, exist_ok=True)
        checkpoint.write_text(
            json.dumps(results, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    return results


def write_enriched_workbook(
    source: Path,
    destination: Path,
    organisations: dict[str, dict],
    results: dict[str, dict],
) -> dict:
    workbook = load_workbook(source)
    for title in ("Organizations", "People"):
        if title in workbook.sheetnames:
            del workbook[title]
    org_sheet = workbook.create_sheet("Organizations")
    org_headers = (
        "abn",
        "business_name",
        "legal_name",
        "website",
        "linkedin_company_url",
        "linkedin_company_name",
        "linkedin_status",
        "website_people_count",
        "linkedin_people_count",
        "website_discovery_status",
    )
    org_sheet.append(org_headers)
    people_sheet = workbook.create_sheet("People")
    people_headers = (
        "abn",
        "business_name",
        "legal_name",
        "person_name",
        "role",
        "linkedin_profile_url",
        "source_type",
        "source_url",
        "confidence",
    )
    people_sheet.append(people_headers)
    people_count = 0
    linkedin_people_count = 0
    for abn, organisation in sorted(
        organisations.items(),
        key=lambda item: (item[1].get("business_name", "").lower(), item[0]),
    ):
        result = results.get(abn) or {
            "status": "not processed",
            "people": [],
            "pages_checked": [],
        }
        people = result.get("people") or []
        linked_people = sum(bool(person.get("linkedin_profile_url")) for person in people)
        org_sheet.append(
            (
                abn,
                organisation.get("business_name", ""),
                organisation.get("legal_name", ""),
                organisation.get("website", ""),
                organisation.get("linkedin_company_url", ""),
                organisation.get("linkedin_profile_name", ""),
                organisation.get("linkedin_status", ""),
                len(people),
                linked_people,
                result.get("status", ""),
            )
        )
        for person in people:
            people_sheet.append(
                (
                    abn,
                    organisation.get("business_name", ""),
                    organisation.get("legal_name", ""),
                    person.get("person_name", ""),
                    person.get("role", ""),
                    person.get("linkedin_profile_url", ""),
                    person.get("source_type", ""),
                    person.get("source_url", ""),
                    person.get("confidence", ""),
                )
            )
            people_count += 1
            linkedin_people_count += bool(person.get("linkedin_profile_url"))
    fill = PatternFill("solid", fgColor="D9FF4F")
    for sheet in (org_sheet, people_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = Font(bold=True, color="14160F")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 24
    for row in org_sheet.iter_rows(min_row=2):
        for column in (4, 5):
            if row[column - 1].value:
                row[column - 1].hyperlink = row[column - 1].value
                row[column - 1].style = "Hyperlink"
    for row in people_sheet.iter_rows(min_row=2):
        for column in (6, 8):
            if row[column - 1].value:
                row[column - 1].hyperlink = row[column - 1].value
                row[column - 1].style = "Hyperlink"
    how = workbook["How to use"]
    start = how.max_row + 2
    notes = (
        "Public people enrichment",
        "No Apify or LinkedIn login was used. LinkedIn itself was not crawled.",
        "People come from official organisation staff, leadership, board, or structured-data pages.",
        "Individual LinkedIn links are included only when an official organisation website publishes them.",
        "Blank results do not prove that an organisation has no employees; public website coverage is incomplete.",
        f"Unique organisations: {len(organisations)}",
        f"Publicly identified people: {people_count}",
        f"People with official-site LinkedIn profile links: {linkedin_people_count}",
    )
    for offset, note in enumerate(notes):
        cell = how.cell(start + offset, 1, note)
        if offset == 0:
            cell.font = Font(bold=True, size=14, color="14160F")
    how.column_dimensions["A"].width = max(how.column_dimensions["A"].width or 0, 130)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return {
        "unique_organisations": len(organisations),
        "public_people": people_count,
        "people_with_linkedin_profile": linkedin_people_count,
        "organisations_with_people": sum(
            bool((results.get(abn) or {}).get("people")) for abn in organisations
        ),
    }
